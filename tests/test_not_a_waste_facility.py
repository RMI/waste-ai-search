"""WP-525: flag facilities that sources say were never disposal sites.

For facilities known only from OSM or Global Plastic Watch, `facility_type` may come back as
"Not a Waste Facility" when a source positively says the site is something else. It is a
contradiction, not a verification: finding nothing leaves facility_type empty.

The failure these tests exist to prevent is the one the ticket names: a closed, capped or
redeveloped landfill read as "not a landfill". Those sites keep emitting methane for decades and
are exactly what WasteMAP tracks, so "it's a park now" must never become "Not a Waste Facility".

By decision on WP-525 the value is promoted like any facility_type - it reaches the standardized
table and the next pass's seed - and it is also always routed to review. Tests below pin both, so
changing either is deliberate.

All offline: agent responses are hand-built payloads fed through the real extraction, resolution,
refresh, standardized-table and workbook code.
"""
from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from waste_ai_search.arbitrate import contradiction_rows, extract_evidence
from waste_ai_search.arbitration import needs_review, resolve
from waste_ai_search.prompt_builder import (
    build_site_prompt,
    needs_contradiction_check,
    requested_attributes,
)
from waste_ai_search.schema import (
    CONTRADICTION_HEADERS,
    NOT_A_WASTE_FACILITY,
    PENDING_UPSTREAM_FACILITY_TYPES,
    map_enum_value,
)


REGULATOR = {
    "source_title": "Mining concessions register",
    "publisher": "Provincial Ministry of Mines",
    "source_type": "Regulator",
    "publication_date": "2023-04-01",
    "url": "https://mines.example.gov/concession/4471",
    "quoted_evidence_short": "Granite quarry, active concession since 1998.",
}

WIKI = {
    "source_title": "Local places",
    "publisher": "Community wiki",
    "source_type": "Wiki or aggregator",
    "publication_date": "",
    "url": "https://wiki.example.org/place/4471",
    "quoted_evidence_short": "The site is a gravel pit.",
}

NEWS = {
    "source_title": "Old tip becomes a park",
    "publisher": "Regional Times",
    "source_type": "News report",
    "publication_date": "2019-06-01",
    "url": "https://news.example.com/tip-park",
    "quoted_evidence_short": "The former landfill was capped in 2009 and reopened as a park.",
}


def unconfirmed_site(**overrides):
    """A facility nothing confirms: known only from OSM, a name and a point."""
    base = {
        "site_id": "4471",
        "internal_facility_id": "4471",
        "site_name": "Landfill",
        "country_iso3": "CAN",
        "latitude": "49.1",
        "longitude": "-122.3",
        "contributing_data_sources": "osm_2022",
        "reference_year": "2022",
    }
    base.update(overrides)
    return base


def attribute(name, value, sources, summary="summary"):
    return {
        "attribute_name": name,
        "value": value,
        "unit": "",
        "value_basis": "Direct",
        "confidence_score": "High",
        "value_date": "2023",
        "evidence_summary": summary,
        "sources": sources,
    }


def arbitrate(site, *attributes):
    """Run one site's payload through extraction and resolution, as run_arbitration does."""
    payload = {"site_id": site["site_id"], "attributes": list(attributes)}
    by_attr, evidence, _sources, _warn = extract_evidence(site, payload, "run", "v", date(2026, 1, 1))
    names = list(requested_attributes(site)) + [n for n in by_attr if n not in requested_attributes(site)]
    resolved = [resolve(site, n, by_attr.get(n, []), reviewed_on=date(2026, 1, 1)) for n in names]
    return resolved, evidence


def row_for(resolved, name):
    return next(r for r in resolved if r["attribute_name"] == name)


def quarry(sources=(REGULATOR,), summary="Active granite quarry."):
    return attribute("facility_type", NOT_A_WASTE_FACILITY, list(sources), summary=summary)


# --- the value ----------------------------------------------------------------------------------
def test_the_value_is_accepted_and_common_phrasings_map_to_it():
    assert map_enum_value("facility_type", NOT_A_WASTE_FACILITY)[0] == NOT_A_WASTE_FACILITY
    for phrasing in ("not waste", "Not a waste site", "not a waste facility"):
        assert map_enum_value("facility_type", phrasing)[0] == NOT_A_WASTE_FACILITY


def test_a_bare_description_is_not_mistaken_for_the_verdict():
    """"granite quarry" in the value field is not a facility type; the verdict has to be explicit."""
    assert map_enum_value("facility_type", "granite quarry")[0] is None


def test_the_ordinary_facility_types_are_unchanged():
    assert map_enum_value("facility_type", "landfill")[0] == "Sanitary Landfill"
    assert map_enum_value("facility_type", "open dump")[0] == "Dumpsite"


# --- criterion 1: a positive contradiction ------------------------------------------------------
def test_a_source_describing_a_quarry_yields_not_a_waste_facility_with_a_clickable_source():
    resolved, _ = arbitrate(unconfirmed_site(), quarry())
    row = row_for(resolved, "facility_type")

    assert row["resolved_value"] == NOT_A_WASTE_FACILITY
    assert row["winning_source_url"].startswith("https://")
    assert row["winning_source_tier"] == "Tier 1"


def test_even_a_tier_1_verdict_is_routed_to_review():
    """Without the override this is an ordinary Tier 1 empty-baseline fill, which auto-validates."""
    resolved, _ = arbitrate(unconfirmed_site(), quarry())
    row = row_for(resolved, "facility_type")

    assert row["validation_status"] == "Needs review"
    assert needs_review(row)
    assert "never signed off automatically" in row["resolution_rule"]


def test_the_override_is_what_prevents_auto_validation():
    """Control: the same Tier 1 source giving an ordinary facility type does auto-validate."""
    resolved, _ = arbitrate(unconfirmed_site(), attribute("facility_type", "Dumpsite", [REGULATOR]))
    assert row_for(resolved, "facility_type")["validation_status"] == "Auto-validated"


# --- a low-tier contradiction -------------------------------------------------------------------
def test_a_low_tier_verdict_still_reaches_review_marked_with_its_tier():
    resolved, _ = arbitrate(unconfirmed_site(), quarry(sources=[WIKI]))
    row = row_for(resolved, "facility_type")

    assert row["resolved_value"] == NOT_A_WASTE_FACILITY
    assert row["winning_source_tier"] == "Tier 4"
    assert needs_review(row)


def test_a_higher_tier_source_wins():
    resolved, _ = arbitrate(unconfirmed_site(), quarry(sources=[WIKI, REGULATOR]))
    assert row_for(resolved, "facility_type")["winning_source_tier"] == "Tier 1"


# --- criterion 2: closure is not a contradiction ------------------------------------------------
def test_a_capped_landfill_keeps_its_type_and_reports_closure():
    """The agent is told a park-on-a-landfill keeps its landfill type, with closure recorded
    through facility_status and closing_year. That is an ordinary fill, not a verdict."""
    resolved, evidence = arbitrate(
        unconfirmed_site(),
        attribute("facility_type", "Sanitary Landfill", [NEWS]),
        attribute("facility_status", "Inactive", [NEWS]),
        attribute("closing_year", "2009", [NEWS]),
    )

    assert row_for(resolved, "facility_type")["resolved_value"] == "Sanitary Landfill"
    assert row_for(resolved, "facility_status")["resolved_value"] == "Inactive"
    assert row_for(resolved, "closing_year")["resolved_value"] == "2009"
    assert contradiction_rows(resolved, evidence) == []


def test_a_capped_landfill_wrongly_called_not_waste_is_caught_by_the_closure_cross_check():
    """If the agent misreads "now a park" anyway, the reviewer is told to look twice."""
    resolved, evidence = arbitrate(
        unconfirmed_site(),
        quarry(sources=[NEWS], summary="Now a public park."),
        attribute("facility_status", "Inactive", [NEWS]),
        attribute("closing_year", "2009", [NEWS]),
    )
    [view] = contradiction_rows(resolved, evidence)

    assert view["closure_also_reported"].startswith("CHECK:")
    assert "facility_status = Inactive" in view["closure_also_reported"]
    assert "closing_year = 2009" in view["closure_also_reported"]


def test_the_guidance_tells_the_agent_closure_is_not_a_contradiction():
    prompt = build_site_prompt(unconfirmed_site(), attributes=requested_attributes(unconfirmed_site()))
    assert "IS STILL A WASTE DISPOSAL" in prompt
    assert "was this ever a waste disposal site?" in prompt
    assert "facility_status = Inactive and closing_year" in prompt


# --- criterion 3: silence is never a negative ---------------------------------------------------
def test_no_contradicting_evidence_leaves_facility_type_empty():
    resolved, evidence = arbitrate(unconfirmed_site())
    row = row_for(resolved, "facility_type")

    assert row["resolved_value"] == ""
    assert row["resolution"] == "Not found"
    assert contradiction_rows(resolved, evidence) == []


def test_a_site_the_agent_could_not_find_gets_no_verdict():
    """An empty attributes list - the agent's "I could not find this facility" - decides nothing."""
    resolved, _ = arbitrate(unconfirmed_site())
    assert all(r["resolved_value"] != NOT_A_WASTE_FACILITY for r in resolved)


# --- criterion 4, as decided on WP-525: promoted, and always reviewed ---------------------------
def test_the_verdict_is_promoted_to_the_standardized_table():
    """Decided on WP-525: flows like any facility_type value. Pinned so reversing it is deliberate."""
    from waste_ai_search.standardized import build_standard_record

    site = unconfirmed_site()
    resolved, _ = arbitrate(site, quarry())
    record, _notes = build_standard_record(site, resolved, sha="deadbeef")

    assert record["facility_type"] == NOT_A_WASTE_FACILITY


def test_the_verdict_is_merged_into_the_next_pass_seed_and_not_offered_again():
    """facility_type is a gap-fill attribute, so the refresh carries it forward. The site then has
    a facility_type, counts as corroborated, and is not offered the verdict a second time."""
    from waste_ai_search.seed_refresh import refresh_sites

    site = unconfirmed_site()
    resolved, _ = arbitrate(site, quarry())
    [refreshed], _stats = refresh_sites([site], resolved, "run")

    assert refreshed["facility_type"] == NOT_A_WASTE_FACILITY
    assert not needs_contradiction_check(refreshed)


def test_the_value_is_declared_as_waiting_on_the_database():
    """The DB enum and chk_facility_type do not carry it yet; the pending list says so explicitly."""
    from waste_ai_search.db_enums import FACILITY_TYPE_VALUES

    assert NOT_A_WASTE_FACILITY in PENDING_UPSTREAM_FACILITY_TYPES
    assert NOT_A_WASTE_FACILITY not in FACILITY_TYPE_VALUES


def test_the_contradictions_view_carries_source_tier_and_quote():
    resolved, evidence = arbitrate(unconfirmed_site(), quarry())
    [view] = contradiction_rows(resolved, evidence)

    assert set(CONTRADICTION_HEADERS) <= set(view)
    assert view["facility_type"] == NOT_A_WASTE_FACILITY
    assert view["winning_source_tier"] == "Tier 1"
    assert view["evidence_summary"] == "Active granite quarry."
    assert "Granite quarry" in view["quoted_evidence_short"]
    assert view["closure_also_reported"] == ""
    assert view["validation_status"] == "Needs review"


def test_the_workbook_has_a_contradictions_tab(tmp_path):
    from waste_ai_search.workbook_io import write_review_workbook

    resolved, evidence = arbitrate(unconfirmed_site(), quarry())
    path = write_review_workbook(
        tmp_path / "r.xlsx",
        run_config=[],
        review_queue=[],
        leads=[],
        parse_warnings=[],
        contradictions=contradiction_rows(resolved, evidence),
    )
    ws = load_workbook(path)["Contradictions"]
    header = [c.value for c in ws[1]]
    assert header == CONTRADICTION_HEADERS
    assert ws.cell(2, header.index("facility_type") + 1).value == NOT_A_WASTE_FACILITY


# --- criterion 5: scope ---------------------------------------------------------------------------
def test_an_osm_only_site_is_offered_the_value():
    assert needs_contradiction_check(unconfirmed_site())
    assert "facility_type" in requested_attributes(unconfirmed_site())


def test_a_gpw_only_site_is_offered_it_too():
    assert needs_contradiction_check(unconfirmed_site(contributing_data_sources="gpw_2021"))


def test_osm_plus_gpw_is_still_unconfirmed():
    """Two Tier 4 detections of one polygon do not establish what it is used for."""
    assert needs_contradiction_check(unconfirmed_site(contributing_data_sources="osm_2022+gpw_2021"))


def test_any_corroborating_baseline_removes_the_offer():
    assert not needs_contradiction_check(unconfirmed_site(facility_status="Active"))
    assert not needs_contradiction_check(unconfirmed_site(annual_incoming_waste_metric_tonnes="12000"))


def test_area_from_the_same_polygon_is_not_corroboration():
    assert needs_contradiction_check(unconfirmed_site(area_square_meters="101999.87"))


def test_a_regulator_backed_site_is_not_offered_it():
    assert not needs_contradiction_check(unconfirmed_site(contributing_data_sources="lmop_2024"))
    assert not needs_contradiction_check(unconfirmed_site(contributing_data_sources="gpw_2021+eprtr_2022"))


def test_a_corroborated_site_prompt_is_not_offered_the_value():
    """Facilities with backing get the four database values and none of the addendum."""
    site = unconfirmed_site(contributing_data_sources="lmop_2024")
    prompt = build_site_prompt(site, attributes=requested_attributes(site))

    assert NOT_A_WASTE_FACILITY not in prompt
    assert "IS STILL A WASTE DISPOSAL" not in prompt
    assert "Allowed values only: Sanitary Landfill, Controlled Dumpsite, Dumpsite, Incineration Facility." in prompt


def test_an_unconfirmed_site_prompt_offers_it_in_guidance_and_definitions():
    prompt = build_site_prompt(unconfirmed_site(), attributes=requested_attributes(unconfirmed_site()))
    definitions = prompt.split("Definitions:")[1]
    assert NOT_A_WASTE_FACILITY in definitions


# --- criterion 6: the premise -------------------------------------------------------------------
def test_the_prompt_no_longer_asserts_the_site_is_a_disposal_facility():
    prompt = build_site_prompt(unconfirmed_site(), attributes=requested_attributes(unconfirmed_site()))
    assert "for this waste disposal site" not in prompt
    assert "recorded as a waste disposal facility" in prompt
