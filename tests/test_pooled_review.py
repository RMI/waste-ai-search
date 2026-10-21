"""Every source filed under a site is reviewed together (build step 5 of the source-matching plan).

After matching, arbitration groups the evidence by the site each source is filed under, from any
search, and makes one decision per (site, attribute): the key ingestion #261 loads by, which stops
on two values for one facility, attribute and year. Every site
column of a review row describes the site the value is filed under; searched_site_id is the search
that found it. These run end to end through run_arbitration, on hand-built source-first responses.
"""
from __future__ import annotations

import csv
import json
import os
import shutil
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from waste_ai_search.arbitrate import run_arbitration
from waste_ai_search.input_loader import write_csv_records
from waste_ai_search.run_context import PipelineConfig
from waste_ai_search.schema import LOCATION_ATTRIBUTE, NOT_A_WASTE_FACILITY
from waste_ai_search.seed_source import seed_headers


def seed(site_id, name, lat, lon, iso3="CAN", **extra):
    return {
        "facility_id": site_id, "site_id": site_id, "internal_facility_id": site_id, "year": "2022",
        "site_name": name, "facility_name": name, "country_iso3": iso3, "iso3c_plus": iso3,
        "latitude": str(lat), "longitude": str(lon), "contributing_data_sources": "osm_2022",
        "reference_year": "2022", **extra,
    }


def source(url, name="", location="", values=(), source_type="Regulator", published="2024-03-01"):
    return {
        "url": url, "source_title": "Page", "publisher": "Publisher", "source_type": source_type,
        "publication_date": published, "facility_name_in_source": name, "facility_name_english": "",
        "stated_location": location, "stated_location_kind": "coordinates" if location else "",
        "values": list(values),
    }


def value(attribute, found, value_date="2024", unit=""):
    return {"attribute_name": attribute, "value": found, "unit": unit, "value_basis": "Direct",
            "confidence_score": "High", "value_date": value_date, "evidence_summary": "Says so."}


def arbitrate(tmp_path, seeds, responses):
    """Arbitrate one run: `responses` maps each site searched to the sources its search cited."""
    run_dir = tmp_path / "run"
    write_csv_records(run_dir / "seed.csv", seeds, seed_headers())
    raw = run_dir / "raw_foundry_responses"
    raw.mkdir(parents=True)
    for site_id, sources in responses.items():
        (raw / f"site_{site_id}.json").write_text(json.dumps({"site_id": site_id, "sources": sources}))
    run_arbitration(PipelineConfig(input_csv=None, run_dir=run_dir, run_id="run", dataset_version="v"))
    return run_dir


def rows(run_dir, name):
    return list(csv.DictReader((run_dir / name).open(encoding="utf-8-sig")))


def queue_key(row):
    """(site, attribute), as ingestion files a validated value."""
    return row["site_id"], row["attribute_name"]


BRADY = seed("13413", "Brady Road Resource Management Facility", 49.7583, -97.2004)
NEARBY = seed("13414", "domestic landfill", 50.10, -97.20)


def test_every_source_filed_under_a_site_is_decided_together(tmp_path):
    """Brady Road's own search found its operator, and so did a search of a site 38 km away whose
    permit places it at Brady Road. One review row, two sources agreeing, all of it Brady Road's."""
    own = source("https://winnipeg.example/brady", "Brady Road Resource Management Facility",
                 values=[value("operator", "City of Winnipeg")])
    permit = source("https://manitoba.example/permit/brady", "Brady Road Landfill", "49.7590, -97.2010",
                    values=[value("operator", "City of Winnipeg")])
    run_dir = arbitrate(tmp_path, [BRADY, NEARBY], {"13413": [own], "13414": [permit]})

    queue = rows(run_dir, "review_queue.csv")
    operator = [row for row in queue if row["attribute_name"] == "operator"]
    assert [(row["site_id"], row["agreeing_source_count"]) for row in operator] == [("13413", "2")]
    [location] = [row for row in queue if row["attribute_name"] == LOCATION_ATTRIBUTE]
    # Every site column is Brady Road's; searched_site_id is the search that found the value.
    assert (location["site_id"], location["site_name"], location["baseline_coordinates"]) == (
        "13413", "Brady Road Resource Management Facility", "49.758300, -97.200400",
    )
    assert (location["searched_site_id"], location["resolved_value"], location["location_basis"]) == (
        "13414", "49.759, -97.201", "coordinates",
    )
    assert location["match"].startswith("Located: by its coordinates, the source places it 0.1 km from this site")
    assert "from the site searched (13414)" in location["match"]
    assert (location["facility_name_in_source"], location["stated_location"]) == ("Brady Road Landfill", "49.7590, -97.2010")
    assert operator[0]["location_basis"] == ""  # only a found_location row says how it was read
    assert not [row for row in queue if row["site_id"] == "13414"]
    assert len({queue_key(row) for row in queue}) == len(queue)

    evidence = rows(run_dir, "evidence.csv")
    assert {(row["site_id"], row["searched_site_id"]) for row in evidence if row["attribute_name"] == "operator"} == {
        ("13413", "13413"), ("13413", "13414"),
    }


def test_values_from_many_searches_never_make_two_rows_for_one_site_and_attribute(tmp_path):
    """A site's own search, a search whose source places it there, and one moved there by name."""
    named = seed("20", "Prairie Green IWMF", 50.0047, -97.1307)
    searched = [seed(sid, "domestic landfill", 50.0 + n / 10, -97.6) for n, sid in enumerate(("21", "22", "23"))]
    status = [value("facility_status", "Active", "2024"), value("facility_status", "Active", "")]
    responses = {
        "20": [source("https://pg.example/own", "Prairie Green IWMF", values=[*status, value("operator", "Waste Connections")])],
        "21": [source("https://pg.example/permit", "Prairie Green Landfill", "50.0050, -97.1310", values=status)],
        "22": [source("https://pg.example/news", "Prairie Green Landfill", values=[*status, value("operator", "WCI")])],
        "23": [source("https://pg.example/news", "Prairie Green Landfill", values=status)],
    }
    run_dir = arbitrate(tmp_path, [named, *searched], responses)
    queue = rows(run_dir, "review_queue.csv")
    keys = Counter(queue_key(row) for row in queue)
    assert keys and max(keys.values()) == 1
    status_rows = [row for row in queue if row["attribute_name"] == "facility_status"]
    assert {row["site_id"] for row in status_rows} == {"20"}
    assert {row["agreeing_source_count"] for row in status_rows} == {"3"}  # own, permit and news
    assert {row["searched_site_id"] for row in queue} <= {"20", "21", "22", "23"}


def test_thunder_bay_coordinates_copied_from_the_prompt_are_no_location_and_the_name_moves_it(tmp_path):
    """Batch 1: a 'domestic landfill' search credited our point to Ontario's registry, and the page
    was Thunder Bay's, 150 km away. Coordinates equal to the prompt's rounded point are no location,
    and the name, which matches Thunder Bay's site strongly, moves the source there."""
    searched = seed("13755", "domestic landfill", 48.843368376131856, -87.42962607259884)
    thunder_bay = seed("18512", "City of Thunder Bay Solid Waste and Recycling Facility", 48.44145, -89.36936,
                       contributing_data_sources="canada_ghgrp_2021")
    page = source("https://www.thunderbay.ca/landfill", "Thunder Bay Solid Waste & Recycling Facility", "48.84, -87.43",
                  values=[value("operator", "City of Thunder Bay")])
    run_dir = arbitrate(tmp_path, [searched, thunder_bay], {"13755": [page]})

    [operator] = [row for row in rows(run_dir, "review_queue.csv") if row["attribute_name"] == "operator"]
    assert (operator["site_id"], operator["searched_site_id"]) == ("18512", "13755")
    assert operator["match"].startswith("Moved by name: the source gives no location, and its name matches this site")
    location = next(row for row in rows(run_dir, "evidence.csv") if row["attribute_name"] == LOCATION_ATTRIBUTE)
    assert (location["location_basis"], location["exclusion_reason"]) == ("none", "No location: copied from the prompt.")


TRUSTED_COORDINATES = "found_latitude@canada_ghgrp_2021; found_longitude@canada_ghgrp_2021"
TRUSTED_NAME = "found_facility_name@canada_ghgrp_2021"


def test_a_trusted_name_or_location_is_judged_on_the_site_a_source_is_filed_under(tmp_path):
    """The source's name and location still match it, but where the MATCHED site's own come from a
    Tier 1-2 source they are recorded, not queued, whatever the site searched."""
    trusted = {**BRADY, "attribute_sources": f"{TRUSTED_COORDINATES}; {TRUSTED_NAME}"}
    permit = source("https://manitoba.example/permit/brady", "Brady Road Landfill", "49.7590, -97.2010",
                    values=[value("found_facility_name", "Brady Road Landfill", ""), value("operator", "City of Winnipeg")])
    run_dir = arbitrate(tmp_path, [trusted, NEARBY], {"13414": [permit]})

    assert [row["attribute_name"] for row in rows(run_dir, "review_queue.csv")] == ["operator"]
    recorded = {row["attribute_name"]: row["exclusion_reason"] for row in rows(run_dir, "evidence.csv")}
    assert recorded[LOCATION_ATTRIBUTE] == "Recorded only: our location comes from a Tier 1-2 source, so it is not reviewed."
    assert recorded["found_facility_name"] == "Recorded only: our name comes from a Tier 1-2 source, so it is not reviewed."

    # The other way round: the site searched is trusted, the site the source is about is not.
    searched = {**NEARBY, "attribute_sources": TRUSTED_COORDINATES}
    run_dir = arbitrate(tmp_path / "again", [BRADY, searched], {"13414": [permit]})
    assert {row["attribute_name"] for row in rows(run_dir, "review_queue.csv")} == {
        LOCATION_ATTRIBUTE, "found_facility_name", "operator",
    }


def test_not_a_waste_facility_stays_only_with_the_site_searched_where_it_was_offered(tmp_path):
    offered = seed("1", "Smith Pit", 50.0, -100.0)
    other = seed("2", "Jones Landfill", 50.3, -100.0)
    verdict = [value("facility_type", "Not a Waste Facility", "")]
    responses = {"1": [
        source("https://county.example/smith-pit", "Smith Pit", values=verdict),
        source("https://county.example/jones", "Jones Landfill", "50.3001, -100.0001", values=verdict),
        source("https://county.example/quarry", "Brown Quarry", "50.6001, -100.0001", values=verdict),
    ]}
    run_dir = arbitrate(tmp_path, [offered, other], responses)

    queue = [row for row in rows(run_dir, "review_queue.csv") if row["attribute_name"] == "facility_type"]
    assert [(row["site_id"], row["resolved_value"]) for row in queue] == [("1", NOT_A_WASTE_FACILITY)]
    leads = [row for row in rows(run_dir, "supplementary_leads.csv") if row["lead_value"] == NOT_A_WASTE_FACILITY]
    assert sorted(row["url"] for row in leads) == ["https://county.example/jones", "https://county.example/quarry"]
    assert all("kept only for the site searched" in row["exclusion_reason"] for row in leads)
    # A place that is no landfill is not a new landfill either.
    assert rows(run_dir, "new_facility_candidates.csv") == []

    # A site with another source is never offered the check, so the verdict is a lead there too.
    corroborated = {**offered, "contributing_data_sources": "osm_2022 + eprtr_2022"}
    run_dir = arbitrate(tmp_path / "again", [corroborated, other], {"1": responses["1"][:1]})
    assert not [row for row in rows(run_dir, "review_queue.csv") if row["attribute_name"] == "facility_type"]
    assert "was not offered it" in rows(run_dir, "supplementary_leads.csv")[0]["exclusion_reason"]


def test_a_match_outside_the_runs_selection_is_listed_and_not_reviewed(tmp_path):
    ours = seed("1", "domestic landfill", 49.2, -97.2)
    across = seed("2", "Pembina Valley Landfill", 48.99, -97.21, "USA")
    page = source("https://nd.example/pembina", "Pembina Valley Landfill", "48.9910, -97.2100",
                  values=[value("operator", "Pembina County")])
    run_dir = arbitrate(tmp_path, [ours, across], {"1": [page]})

    assert not [row for row in rows(run_dir, "review_queue.csv") if row["attribute_name"] == "operator"]
    [listed] = rows(run_dir, "outside_run_matches.csv")
    assert (listed["matched_site_id"], listed["country_iso3"], listed["searched_site_id"]) == ("2", "USA", "1")
    config = {row[0]: row[1] for row in load_workbook(run_dir / "run_review.xlsx")["Run_Config"].values}
    assert (config["match_located"], config["match_outside_run"]) == (1, 1)


def test_the_lists_for_ingestion_are_written_at_arbitration(tmp_path):
    cowpen = [seed("18572", "Cowpen Bewley Landfill Site", 54.6136, -1.24064, "GBR"),
              seed("19041", "Cowpen Bewley Landfill Site", 54.613634, -1.240635, "GBR")]
    register = "https://environment.data.gov.uk/public-register/EPR-RP3531DV"
    new = source("https://council.example/new-tip", "Hartlepool Road Landfill", "54.7001, -1.3001",
                 values=[value("operator", "Hartlepool Council")])
    # Both searches cited the register, with no location; a source with no value and no location has
    # nothing to file, so these give the operator.
    listed = source(register, "Cowpen Bewley Landfill Site", values=[value("operator", "Augean")])
    responses = {"18572": [listed, new], "19041": [listed]}
    run_dir = arbitrate(tmp_path, cowpen, responses)

    [pair] = rows(run_dir, "possible_duplicates.csv")
    assert (pair["site_id_a"], pair["site_id_b"], pair["distance_km"], pair["searched_site_ids"]) == ("18572", "19041", "0.00", "18572; 19041")
    [candidate] = rows(run_dir, "new_facility_candidates.csv")
    assert (candidate["source_urls"], candidate["facility_names_in_source"], candidate["stated_locations"]) == (
        "https://council.example/new-tip", "Hartlepool Road Landfill", "54.7001, -1.3001",
    )
    assert (candidate["nearest_site_id"], candidate["searched_site_ids"], candidate["location_basis"]) == ("19041", "18572", "coordinates")
    # Its values belong to no site of ours, so none is reviewed; the register's stay with each search.
    operators = {(row["site_id"], row["resolved_value"]) for row in rows(run_dir, "review_queue.csv") if row["attribute_name"] == "operator"}
    assert operators == {("18572", "Augean"), ("19041", "Augean")}
    config = {row[0]: row[1] for row in load_workbook(run_dir / "run_review.xlsx")["Run_Config"].values}
    assert (config["new_facility_candidates"], config["possible_duplicates"], config["match_new_facility"]) == (1, 1, 1)


# --- the gas follow-up --------------------------------------------------------------------------
class _GasClient:
    """Pass 1 finds gas collection on a page about a nearby site; the follow-up gives a GCCS value."""

    def __init__(self, located_at):
        self.located_at = located_at
        self.prompts = []

    def search_site(self, site, prompt):  # noqa: ARG002
        self.prompts.append(prompt)
        page = source("https://reg.example/permit", "Second Landfill", self.located_at, published="2024-01-01")
        if "gccs_ch4_flared_metric_tonnes" in prompt:
            page["values"] = [value("gccs_ch4_flared_metric_tonnes", "1200", unit="metric tonnes")]
        else:
            page["values"] = [value("has_landfill_gas_collection", "Yes")]
        return {"sources": [page]}


@pytest.mark.parametrize("country, followed_up", [("PHL", True), ("USA", False)])
def test_the_follow_up_unlocks_the_matched_site_only_within_the_runs_selection(tmp_path, monkeypatch, country, followed_up):
    from waste_ai_search import arbitrate as ar
    from waste_ai_search import cli, search as se

    seed_csv = tmp_path / "seed.csv"
    seed_csv.write_text(
        "site_id,internal_facility_id,site_name,country_iso3,latitude,longitude,"
        "is_location_exact,has_landfill_gas_collection,reference_year\n"
        "1,1,First Landfill,PHL,6.5,3.3,TRUE,,2022\n"
        f"2,2,Second Landfill,{country},6.6,3.3,TRUE,,2022\n",
        encoding="utf-8",
    )
    client = _GasClient("6.6001, 3.3001")
    monkeypatch.setattr(se, "get_client", lambda config: client)
    monkeypatch.setattr(ar, "write_review_workbook", lambda path, **kw: path)
    cli.run_everything(PipelineConfig(input_csv=seed_csv, run_dir=tmp_path / "run", run_id="t", site_delay_seconds=0, limit=1))

    assert len(client.prompts) == (2 if followed_up else 1)
    if followed_up:
        assert "Second Landfill" in client.prompts[1] and "gccs_ch4_flared_metric_tonnes" in client.prompts[1]


# --- batch 1, replayed --------------------------------------------------------------------------
REPLAY_RUN = os.environ.get("WASTE_AI_SEARCH_REPLAY_RUN", "")


@pytest.mark.skipif(not REPLAY_RUN, reason="set WASTE_AI_SEARCH_REPLAY_RUN to a batch-1 run folder to replay it")
def test_a_batch_1_replay_never_has_two_rows_for_one_site_and_attribute(tmp_path):
    """Batch 1's cached responses, re-arbitrated offline: a plumbing check. The run folder holds
    seed.csv and raw_foundry_responses/, and link_check.json if it was checked. It is copied, so
    the folder itself is left as it is."""
    source_dir = Path(REPLAY_RUN)
    run_dir = tmp_path / "batch1"
    run_dir.mkdir()
    shutil.copy(source_dir / "seed.csv", run_dir / "seed.csv")
    if (source_dir / "link_check.json").exists():
        shutil.copy(source_dir / "link_check.json", run_dir / "link_check.json")
    (run_dir / "raw_foundry_responses").symlink_to((source_dir / "raw_foundry_responses").resolve())
    run_arbitration(PipelineConfig(input_csv=None, run_dir=run_dir, run_id="batch1", dataset_version="v"))

    queue = rows(run_dir, "review_queue.csv")
    assert queue
    keys = Counter(queue_key(row) for row in queue)
    assert [key for key, count in keys.items() if count > 1] == []
    assert all(row["searched_site_id"] and row["match"] for row in queue)
