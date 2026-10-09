"""The workbook is the only file an SME opens, so its affordances must match the instructions."""
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from waste_ai_search.schema import DEFINITION_VALUES, REVIEW_QUEUE_HEADERS
from waste_ai_search.workbook_io import write_review_workbook


def build(tmp_path, queue_rows=None):
    path = write_review_workbook(
        tmp_path / "r_review.xlsx",
        run_config=[{"setting": "run_id", "value": "r"}],
        review_queue=queue_rows
        if queue_rows is not None
        else [{h: "" for h in REVIEW_QUEUE_HEADERS} | {"site_id": "1", "resolution": "Filled empty baseline"}],
        leads=[],
        parse_warnings=[],
    )
    return load_workbook(path)


def test_expected_tabs_exist(tmp_path):
    wb = build(tmp_path)
    assert wb.sheetnames == [
        "ReadMe", "Run_Config", "Review_Queue", "Leads", "Parse_Warnings", "Definitions",
    ]


def test_only_the_review_decision_columns_have_dropdowns(tmp_path):
    """A dropdown on a pipeline-set column reads as permission to change it."""
    ws = build(tmp_path)["Review_Queue"]
    headers = {ws.cell(1, c).value: get_column_letter(c) for c in range(1, ws.max_column + 1)}

    validated_columns = set()
    for dv in ws.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            validated_columns.add("".join(ch for ch in rng.split(":")[0] if ch.isalpha()))

    assert validated_columns == {headers["validation_status"], headers["rejection_reason"]}  # no facility_type row here
    assert headers["resolution"] not in validated_columns
    assert headers["winning_source_tier"] not in validated_columns


def test_the_dropdown_offers_only_reviewer_choosable_outcomes(tmp_path):
    wb = build(tmp_path)
    ws = wb["Review_Queue"]
    defs = wb["Definitions"]
    columns = {defs.cell(1, c).value: c for c in range(1, defs.max_column + 1)}

    queue = {ws.cell(1, c).value: get_column_letter(c) for c in range(1, ws.max_column + 1)}
    dv = next(d for d in ws.data_validations.dataValidation if str(d.sqref).startswith(queue["validation_status"]))
    letter = get_column_letter(columns["review_decision"])
    assert f"${letter}$" in dv.formula1
    assert DEFINITION_VALUES["review_decision"] == ["Validated", "Rejected"]


def test_definitions_tab_publishes_the_tier_scale(tmp_path):
    defs = build(tmp_path)["Definitions"]
    rows = list(defs.values)
    index = rows[0].index("source_tier")
    assert [r[index] for r in rows[1:] if r[index]] == ["Tier 1", "Tier 2", "Tier 3", "Tier 4", "Tier 5"]


def test_readme_tab_documents_every_tab_and_the_csv_sidecars(tmp_path):
    wb = build(tmp_path)
    described = {row[0] for row in list(wb["ReadMe"].values)[1:]}
    for tab in ("Review_Queue", "Leads", "Parse_Warnings", "Definitions"):
        assert tab in described
    assert any(str(name).startswith("(file)") for name in described)


def test_a_rejection_reason_is_picked_from_a_fixed_list(tmp_path):
    """Rejections are counted by cause to find where the AI errs, so the reason is a dropdown."""
    wb = build(tmp_path)
    ws = wb["Review_Queue"]
    defs = wb["Definitions"]
    columns = {defs.cell(1, c).value: c for c in range(1, defs.max_column + 1)}
    queue = {ws.cell(1, c).value: get_column_letter(c) for c in range(1, ws.max_column + 1)}

    dv = next(d for d in ws.data_validations.dataValidation if str(d.sqref).startswith(queue["rejection_reason"]))
    letter = get_column_letter(columns["rejection_reason"])
    assert f"${letter}$2:${letter}$6" in dv.formula1
    offered = [defs.cell(r, columns["rejection_reason"]).value for r in range(2, 7)]
    assert offered == ["Wrong facility", "Not in source", "Wrong value", "Outdated", "Other"]
    assert REVIEW_QUEUE_HEADERS.index("rejection_reason") == REVIEW_QUEUE_HEADERS.index("validation_status") + 1


def test_baseline_coordinates_paste_straight_into_google_maps():
    """SMEs copy the seed location into Google Maps' search box, which takes "lat, lon"."""
    from waste_ai_search.arbitrate import baseline_coordinates

    assert baseline_coordinates({"latitude": "45.51234", "longitude": "-73.55432"}) == "45.512340, -73.554320"
    assert baseline_coordinates({"latitude": "", "longitude": "-73.5"}) == ""
    assert REVIEW_QUEUE_HEADERS.index("baseline_coordinates") == REVIEW_QUEUE_HEADERS.index("country_iso3") + 1


def test_corrected_value_offers_the_values_of_the_attribute_on_its_row(tmp_path):
    """Site 2529: an SME picks "Transfer Station" without overwriting the AI's own answer."""
    blank = {h: "" for h in REVIEW_QUEUE_HEADERS}
    rows = [
        ("facility_type", "Not a Waste Facility"),   # row 2
        ("opening_year", "1990"),                    # row 3: free text
        ("has_flare", "TRUE"),                       # row 4
        ("cover_types", "clay cover"),               # row 5
        ("facility_type", "Dumpsite"),               # row 6
    ]
    wb = build(tmp_path, [blank | {"site_id": "1", "attribute_name": a, "resolved_value": v} for a, v in rows])
    ws = wb["Review_Queue"]
    defs = wb["Definitions"]
    target = {ws.cell(1, c).value: get_column_letter(c) for c in range(1, ws.max_column + 1)}["corrected_value"]
    def_col = {defs.cell(1, c).value: get_column_letter(c) for c in range(1, defs.max_column + 1)}
    by_list = {d.formula1: d for d in ws.data_validations.dataValidation if target in str(d.sqref)}

    ft = by_list[next(f for f in by_list if f"${def_col['facility_type']}$" in f)]
    assert str(ft.sqref) == f"{target}2 {target}6" and ft.showErrorMessage
    flare = by_list[next(f for f in by_list if f"${def_col['true_false']}$" in f)]
    assert str(flare.sqref) == f"{target}4" and flare.showErrorMessage
    cover = by_list[next(f for f in by_list if f"${def_col['cover_type']}$" in f)]
    assert str(cover.sqref) == f"{target}5" and not cover.showErrorMessage  # combinations may be typed
    assert not any(f"{target}3" in str(d.sqref).split() for d in by_list.values())  # opening_year: free text
    assert "Transfer Station" in DEFINITION_VALUES["facility_type"]
    assert REVIEW_QUEUE_HEADERS.index("corrected_value") == REVIEW_QUEUE_HEADERS.index("rejection_reason") + 1
