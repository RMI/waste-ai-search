"""Re-keying a pre-wastemap_id run (batch 1) from internal_facility_id to wastemap_id."""
from __future__ import annotations

import csv
import gzip
import hashlib
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

import waste_ai_search.db as db
from waste_ai_search.input_loader import write_csv_records
from waste_ai_search.relabel import (
    MAPPED,
    MERGED,
    PARTLY_LINKED,
    SNAPSHOT_CROSSWALK,
    SPLIT,
    UNMAPPED,
    Crosswalk,
    RelabelError,
    choose_consolidation_run,
    map_old_ids,
    read_crosswalk_csv,
    read_crosswalk_db,
    relabel_run,
)
from waste_ai_search.schema import (
    PARSE_WARNING_HEADERS,
    RESOLVED_HEADERS,
    REVIEW_QUEUE_HEADERS,
    SUPPLEMENTARY_LEADS_HEADERS,
)
from waste_ai_search.workbook_io import write_review_workbook

RUN_ID = "batchX"
# Old ID -> the source records the run's own crosswalk linked to it.
SNAPSHOT = {
    "1": [("osm_2022", "A")],
    "2": [("eprtr_2022", "B"), ("osm_2022", "C")],
    "3": [("osm_2022", "D")],
}
# Today's crosswalk, with every record linked and each old ID's records on one facility.
CLEAN = {("osm_2022", "A"): "101", ("eprtr_2022", "B"): "102", ("osm_2022", "C"): "102", ("osm_2022", "D"): "103"}
CLEAN_MAP = {"1": "101", "2": "102", "3": "103"}


def crosswalk(records: dict) -> Crosswalk:
    return Crosswalk({record: {wastemap_id} for record, wastemap_id in records.items()}, "test crosswalk")


def snapshot_sets() -> dict:
    return {old_id: set(records) for old_id, records in SNAPSHOT.items()}


def resolved_rows(ids=("1", "2", "3")) -> list[dict]:
    blank = {header: "" for header in RESOLVED_HEADERS}
    return [
        blank | {
            "site_id": site_id,
            "internal_facility_id": site_id,
            # Quotes, a comma and a line break, so the rewrite must quote exactly as the search did.
            "site_name": f'Site {site_id}, "North"\nCell',
            "attribute_name": "facility_status",
            "resolved_value": "Inactive",
            "winning_evidence_id": f"WAIEV-{site_id}-001-001",
            "researcher_notes": "Source says 'closed in 2017'.",
        }
        for site_id in ids
    ]


def make_run(tmp_path, *, ids=("1", "2", "3"), queue_ids=("1", "3"), extra_queue_column=None, numeric_ids=False):
    run_dir = tmp_path / RUN_ID
    run_dir.mkdir()
    write_csv_records(run_dir / "resolved.csv", resolved_rows(ids), RESOLVED_HEADERS)

    queue_headers = REVIEW_QUEUE_HEADERS + ([extra_queue_column] if extra_queue_column else [])
    blank = {header: "" for header in queue_headers}
    to_cell = int if numeric_ids else str
    write_review_workbook(
        run_dir / f"{RUN_ID}_review.xlsx",
        run_config=[{"setting": "run_id", "value": RUN_ID}, {"setting": "sites_arbitrated", "value": 3}],
        review_queue=[
            blank | {"site_id": to_cell(site_id), "attribute_name": "opening_year", "validation_status": "Validated"}
            for site_id in queue_ids
        ],
        leads=[{"run_id": RUN_ID, "site_id": to_cell("2"), "lead_value": "1990"}],
        parse_warnings=[{"run_id": RUN_ID, "site_id": to_cell("3"), "warning": "unit"}],
    )
    if extra_queue_column:
        workbook = load_workbook(run_dir / f"{RUN_ID}_review.xlsx")
        sheet = workbook["Review_Queue"]
        sheet.cell(1, len(queue_headers), extra_queue_column)
        workbook.save(run_dir / f"{RUN_ID}_review.xlsx")
    # An SME's decision date, typed into Excel as a date.
    workbook = load_workbook(run_dir / f"{RUN_ID}_review.xlsx")
    sheet = workbook["Review_Queue"]
    sheet.cell(2, REVIEW_QUEUE_HEADERS.index("reviewed_date") + 1, datetime(2026, 10, 9))
    workbook.save(run_dir / f"{RUN_ID}_review.xlsx")

    snapshot = run_dir / SNAPSHOT_CROSSWALK
    snapshot.parent.mkdir()
    with gzip.open(snapshot, "wt", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["internal_facility_id", "data_source", "facility_id", "entity_linkage_strategy"])
        for old_id, records in SNAPSHOT.items():
            for data_source, facility_id in records:
                writer.writerow([old_id, data_source, facility_id, "legacy_replica_v1"])
    return run_dir


def digest(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sheet_rows(path, name) -> list[tuple]:
    workbook = load_workbook(path, read_only=True)
    try:
        return [tuple(row) for row in workbook[name].iter_rows(values_only=True)]
    finally:
        workbook.close()


# --- mapping ------------------------------------------------------------------------------------------

def test_ids_whose_records_are_all_on_one_facility_map_one_to_one():
    mappings = map_old_ids(["1", "2", "3"], snapshot_sets(), crosswalk(CLEAN))

    assert {old_id: (m.outcome, m.wastemap_id) for old_id, m in mappings.items()} == {
        "1": (MAPPED, "101"), "2": (MAPPED, "102"), "3": (MAPPED, "103"),
    }
    assert (mappings["2"].source_records, mappings["2"].linked_records) == (2, 2)


def test_an_id_still_maps_when_some_of_its_records_are_no_longer_linked():
    records = {record: wastemap_id for record, wastemap_id in CLEAN.items() if record != ("osm_2022", "C")}

    mapping = map_old_ids(["2"], snapshot_sets(), crosswalk(records))["2"]

    assert (mapping.outcome, mapping.wastemap_id, mapping.linked_records) == (PARTLY_LINKED, "102", 1)
    assert "osm_2022:C" in mapping.detail


def test_an_id_none_of_whose_records_is_linked_is_unmapped():
    records = {record: wastemap_id for record, wastemap_id in CLEAN.items() if record[1] not in "BC"}

    mapping = map_old_ids(["2"], snapshot_sets(), crosswalk(records))["2"]

    assert (mapping.outcome, mapping.wastemap_id, mapping.candidates) == (UNMAPPED, "", [])
    assert "eprtr_2022:B, osm_2022:C" in mapping.detail


def test_an_id_missing_from_the_runs_own_crosswalk_is_unmapped():
    mapping = map_old_ids(["9"], snapshot_sets(), crosswalk(CLEAN))["9"]

    assert (mapping.outcome, mapping.detail) == (UNMAPPED, "not in the run's own crosswalk")


def test_an_id_whose_records_now_belong_to_two_facilities_is_split():
    mapping = map_old_ids(["2"], snapshot_sets(), crosswalk(CLEAN | {("osm_2022", "C"): "104"}))["2"]

    assert (mapping.outcome, mapping.wastemap_id, mapping.candidates) == (SPLIT, "", ["102", "104"])
    assert "102 <- eprtr_2022:B" in mapping.detail and "104 <- osm_2022:C" in mapping.detail


def test_two_ids_that_now_share_a_facility_are_both_merged():
    mappings = map_old_ids(["1", "2", "3"], snapshot_sets(), crosswalk(CLEAN | {("osm_2022", "D"): "101"}))

    assert {old_id: (m.outcome, m.wastemap_id) for old_id, m in mappings.items()} == {
        "1": (MERGED, ""), "2": (MAPPED, "102"), "3": (MERGED, ""),
    }
    assert "also the facility of old ID(s) 3" in mappings["1"].detail


def test_a_merge_with_an_id_the_run_never_searched_is_no_merge():
    """Only the run's own IDs can collide; a facility that absorbed an unsearched one is fine."""
    mappings = map_old_ids(["1", "2"], snapshot_sets(), crosswalk(CLEAN | {("osm_2022", "D"): "101"}))

    assert mappings["1"].outcome == MAPPED


# --- relabelling a run --------------------------------------------------------------------------------

def test_relabel_rewrites_the_ids_and_nothing_else(tmp_path):
    run_dir = make_run(tmp_path)
    originals = {path.name: digest(path) for path in run_dir.iterdir() if path.is_file()}

    result = relabel_run(run_dir, crosswalk(CLEAN), today=date(2026, 10, 19))

    out_dir = run_dir / "wastemap_id"
    assert result.written == [out_dir / "resolved.csv", out_dir / f"{RUN_ID}_review.xlsx"]
    assert {path.name: digest(path) for path in run_dir.iterdir() if path.is_file()} == originals

    # resolved.csv: byte for byte what the search would have written with the new IDs.
    expected = tmp_path / "expected.csv"
    rows = resolved_rows()
    for row in rows:
        row["site_id"] = row["internal_facility_id"] = CLEAN_MAP[row["site_id"]]
    write_csv_records(expected, rows, RESOLVED_HEADERS)
    assert (out_dir / "resolved.csv").read_bytes() == expected.read_bytes()
    assert "WAIEV-1-001-001" in (out_dir / "resolved.csv").read_text()  # evidence keys still point into evidence.csv

    copy = out_dir / f"{RUN_ID}_review.xlsx"
    original = run_dir / f"{RUN_ID}_review.xlsx"
    for name, column in (("Review_Queue", 0), ("Leads", 1), ("Parse_Warnings", 1)):
        before, after = sheet_rows(original, name), sheet_rows(copy, name)
        assert [row[column] for row in after[1:]] == [CLEAN_MAP[row[column]] for row in before[1:]]
        assert [row[:column] + row[column + 1:] for row in after] == [row[:column] + row[column + 1:] for row in before]
    assert sheet_rows(copy, "Review_Queue")[1][REVIEW_QUEUE_HEADERS.index("reviewed_date")] == datetime(2026, 10, 9)

    settings = sheet_rows(copy, "Run_Config")
    assert settings[:3] == sheet_rows(original, "Run_Config")
    assert settings[3] == ("site_id_key", "wastemap_id")
    assert settings[4][0] == "site_id_relabelled" and settings[4][1].startswith("2026-10-19: ")
    assert "test crosswalk" in settings[4][1]
    # The SMEs' dropdowns survive.
    assert len(load_workbook(copy)["Review_Queue"].data_validations.dataValidation) >= 2

    with (out_dir / "relabel_report.csv").open(newline="", encoding="utf-8") as handle:
        report = list(csv.DictReader(handle))
    assert {row["old_internal_facility_id"]: row["wastemap_id"] for row in report} == CLEAN_MAP
    assert report[0]["appears_in"] == "resolved.csv: 1; Review_Queue: 1"


@pytest.mark.parametrize(
    "records, outcome, problem_ids",
    [
        ({record: wastemap_id for record, wastemap_id in CLEAN.items() if record[1] != "A"}, UNMAPPED, {"1"}),
        (CLEAN | {("osm_2022", "C"): "104"}, SPLIT, {"2"}),
        (CLEAN | {("osm_2022", "D"): "101"}, MERGED, {"1", "3"}),
    ],
)
def test_relabel_stops_and_writes_only_the_report(tmp_path, records, outcome, problem_ids):
    run_dir = make_run(tmp_path)

    result = relabel_run(run_dir, crosswalk(records))

    assert {m.old_id for m in result.stopped} == problem_ids
    assert result.written == []
    assert sorted(path.name for path in (run_dir / "wastemap_id").iterdir()) == ["relabel_report.csv"]
    with result.report.open(newline="", encoding="utf-8") as handle:
        report = list(csv.DictReader(handle))
    assert [row["outcome"] for row in report[: len(problem_ids)]] == [outcome] * len(problem_ids)


def test_a_dry_run_writes_only_the_report(tmp_path):
    run_dir = make_run(tmp_path)

    result = relabel_run(run_dir, crosswalk(CLEAN), dry_run=True)

    assert result.written == [] and not result.stopped
    assert sorted(path.name for path in (run_dir / "wastemap_id").iterdir()) == ["relabel_report.csv"]


def test_the_smes_downloaded_workbook_is_written_under_its_runs_name(tmp_path):
    run_dir = make_run(tmp_path)
    downloads = tmp_path / "Downloads"
    downloads.mkdir()
    sme_copy = downloads / f"{RUN_ID}_review (2).xlsx"
    sme_copy.write_bytes((run_dir / f"{RUN_ID}_review.xlsx").read_bytes())

    result = relabel_run(run_dir, crosswalk(CLEAN), out_dir=tmp_path / "out", workbook_path=sme_copy)

    assert [path.name for path in result.written] == ["resolved.csv", f"{RUN_ID}_review.xlsx"]


def test_numeric_site_id_cells_stay_numeric(tmp_path):
    """Excel can turn a typed ID into a number; the copy keeps the cell's type."""
    run_dir = make_run(tmp_path, numeric_ids=True)

    relabel_run(run_dir, crosswalk(CLEAN))

    rows = sheet_rows(run_dir / "wastemap_id" / f"{RUN_ID}_review.xlsx", "Review_Queue")
    assert [row[0] for row in rows[1:]] == [101, 103]


def test_a_workbook_with_corrected_value_date_is_accepted(tmp_path):
    """The raw load takes corrected_value_date when a workbook has it, and loads it empty when not."""
    run_dir = make_run(tmp_path, extra_queue_column="corrected_value_date")

    assert relabel_run(run_dir, crosswalk(CLEAN)).written


def test_a_column_the_raw_load_would_refuse_stops_it(tmp_path):
    run_dir = make_run(tmp_path, extra_queue_column="sme_comment")

    with pytest.raises(RelabelError, match=r"unexpected \['sme_comment'\]"):
        relabel_run(run_dir, crosswalk(CLEAN))
    assert not (run_dir / "wastemap_id").exists()


def test_a_relabelled_workbook_is_not_relabelled_again(tmp_path):
    """Its IDs are wastemap_ids now; reading them as old IDs would silently re-key them wrongly."""
    run_dir = make_run(tmp_path)
    relabel_run(run_dir, crosswalk(CLEAN))

    with pytest.raises(RelabelError, match="has been relabelled"):
        relabel_run(
            run_dir,
            crosswalk(CLEAN),
            out_dir=tmp_path / "again",
            workbook_path=run_dir / "wastemap_id" / f"{RUN_ID}_review.xlsx",
        )


def test_copies_are_never_written_over(tmp_path):
    run_dir = make_run(tmp_path)
    relabel_run(run_dir, crosswalk(CLEAN))

    with pytest.raises(RelabelError, match="already exists"):
        relabel_run(run_dir, crosswalk(CLEAN))
    with pytest.raises(RelabelError, match="folder of their own"):
        relabel_run(run_dir, crosswalk(CLEAN), out_dir=run_dir)


def test_a_copy_that_changed_more_than_its_ids_is_not_kept(tmp_path, monkeypatch):
    import waste_ai_search.relabel as relabel

    write_workbook = relabel.write_workbook

    def careless(workbook, path, id_map, provenance):
        workbook["Review_Queue"].cell(2, REVIEW_QUEUE_HEADERS.index("validation_status") + 1, "Rejected")
        write_workbook(workbook, path, id_map, provenance)

    monkeypatch.setattr(relabel, "write_workbook", careless)
    run_dir = make_run(tmp_path)

    with pytest.raises(RelabelError, match="Review_Queue row 2 changed beyond its site_id"):
        relabel_run(run_dir, crosswalk(CLEAN))
    assert sorted(path.name for path in (run_dir / "wastemap_id").iterdir()) == ["relabel_report.csv"]


def test_resolved_rows_naming_two_facilities_stop_it(tmp_path):
    run_dir = make_run(tmp_path)
    rows = resolved_rows()
    rows[1]["internal_facility_id"] = "3"
    write_csv_records(run_dir / "resolved.csv", rows, RESOLVED_HEADERS)

    with pytest.raises(RelabelError, match="site_id and internal_facility_id differ"):
        relabel_run(run_dir, crosswalk(CLEAN))


# --- reading today's crosswalk ------------------------------------------------------------------------

def write_crosswalk_csv(path, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return path


def test_a_crosswalk_export_is_read_whatever_its_id_format(tmp_path):
    path = write_crosswalk_csv(
        tmp_path / "crosswalk.csv",
        [
            {"wastemap_id": "101.0", "data_source": "osm_2022", "facility_id": "A", "entity_linkage_run_id": "el-1"},
            {"wastemap_id": "102", "data_source": "eprtr_2022", "facility_id": " B ", "entity_linkage_run_id": "el-1"},
        ],
    )

    loaded = read_crosswalk_csv(path)

    assert loaded.wastemap_ids == {("osm_2022", "A"): {"101"}, ("eprtr_2022", "B"): {"102"}}
    assert "el-1" in loaded.description


def test_a_crosswalk_export_holding_two_runs_is_refused(tmp_path):
    path = write_crosswalk_csv(
        tmp_path / "crosswalk.csv",
        [
            {"wastemap_id": "101", "data_source": "osm_2022", "facility_id": "A", "entity_linkage_run_id": "el-1"},
            {"wastemap_id": "205", "data_source": "osm_2022", "facility_id": "A", "entity_linkage_run_id": "el-2"},
        ],
    )

    with pytest.raises(RelabelError, match="2 entity linkage runs"):
        read_crosswalk_csv(path)


RUNS = [
    {"consolidation_run_id": "c-3", "consolidation_strategy": "exact_name_v1", "entity_linkage_strategy": "exact_name_v1",
     "entity_linkage_run_id": "el-3", "created": datetime(2026, 10, 18)},
    {"consolidation_run_id": "c-2", "consolidation_strategy": "legacy_replica_v1", "entity_linkage_strategy": "legacy_replica_v1",
     "entity_linkage_run_id": "el-2", "created": datetime(2026, 10, 17)},
    {"consolidation_run_id": "c-1", "consolidation_strategy": "legacy_replica_v1", "entity_linkage_strategy": "legacy_replica_v1",
     "entity_linkage_run_id": "el-1", "created": datetime(2026, 10, 1)},
]


def test_the_newest_legacy_replica_run_is_read_by_default():
    """The same run a seed reads, not the newest run overall, which can be an experiment."""
    assert choose_consolidation_run(RUNS, environ={})["entity_linkage_run_id"] == "el-2"


def test_a_pinned_consolidation_run_is_read():
    assert choose_consolidation_run(RUNS, "c-1", environ={})["entity_linkage_run_id"] == "el-1"
    assert choose_consolidation_run(RUNS, environ={"CONSOLIDATION_RUN_ID": "c-1"})["entity_linkage_run_id"] == "el-1"
    with pytest.raises(RelabelError, match="not in consolidation"):
        choose_consolidation_run(RUNS, "c-9", environ={})


def test_the_database_crosswalk_is_the_run_the_consolidation_consumed(monkeypatch):
    calls = []

    def fetch_all(sql, params=None, config=None):  # noqa: ARG001
        calls.append((sql, params))
        if "consolidated_facility" in sql:
            return RUNS
        return [{"wastemap_id": 101, "data_source": "osm_2022", "facility_id": "A"}]

    monkeypatch.setattr(db, "fetch_all", fetch_all)
    monkeypatch.delenv("CONSOLIDATION_RUN_ID", raising=False)
    monkeypatch.delenv("CONSOLIDATION_STRATEGY", raising=False)
    monkeypatch.delenv("ENTITY_LINKAGE_STRATEGY", raising=False)

    loaded = read_crosswalk_db()

    assert calls[1][1] == ("el-2",)
    assert loaded.wastemap_ids == {("osm_2022", "A"): {"101"}}
    assert "consolidation run c-2" in loaded.description
