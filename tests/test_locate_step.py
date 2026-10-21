"""The locate step: it runs in the search phase after the link check, and only it uses the network.

Every source's stated location is located once and the answer kept in source_locations.json,
which arbitration reads, so arbitration stays offline. `locate --run-id` backfills an older run,
as `check-links` does, and `--no-locate` skips it. No test here reaches a real geocoder.
"""
from __future__ import annotations

import csv
import json

from openpyxl import load_workbook

from waste_ai_search import cli, search
from waste_ai_search.arbitrate import run_arbitration
from waste_ai_search.locate import NO_GEOCODER, NOT_LOCATED, ROOFTOP, SOURCE_LOCATIONS_FILE, GeocodeMatch, load_locations
from waste_ai_search.run_context import PipelineConfig


PERMIT = {
    "source_title": "Waste licence W0123",
    "publisher": "Environmental Protection Agency",
    "source_type": "Permit or registry",
    "publication_date": "2022-03-01",
}


def response(*sources):
    return {"site_id": "1", "sources": list(sources), "unverified_leads": [], "search_notes": "Searched."}


def source(url, stated, kind, **extra):
    return {**PERMIT, "url": url, "stated_location": stated, "stated_location_kind": kind, "values": [], **extra}


ADDRESS = source("https://epa.example.ie/licence/W0123", "Ballynagran, Coolbeg, Co. Wicklow", "street address")
GRID = source("https://epa.example.ie/licence/W0124", "T 285 935", "grid reference")


class FakeGeocoder:
    name = "fake"

    def __init__(self):
        self.calls = []

    def geocode(self, address, iso2):
        self.calls.append((address, iso2))
        return GeocodeMatch(52.9500, -6.0900, ROOFTOP, "Address, Rooftop, High, Good")


class Client:
    def __init__(self, payload):
        self.payload = payload

    def search_site(self, site, prompt):  # noqa: ARG002
        return self.payload


def seed(tmp_path):
    path = tmp_path / "seed.csv"
    path.write_text(
        "site_id,internal_facility_id,site_name,country_iso3,latitude,longitude,contributing_data_sources,reference_year\n"
        "1,1,Ballynagran Landfill,IRL,52.95,-6.10,osm_2022,2022\n",
        encoding="utf-8",
    )
    return path


def run_search(tmp_path, monkeypatch, geocoder=None, **config):
    monkeypatch.setattr(search, "get_client", lambda _config: Client(response(ADDRESS, GRID)))
    monkeypatch.setattr(search, "make_geocoder", lambda name: geocoder)
    run_config = PipelineConfig(
        input_csv=seed(tmp_path), run_dir=tmp_path / "run", run_id="run", site_delay_seconds=0, **config
    )
    search.run_search(run_config)
    return run_config


def test_the_search_locates_each_stated_location_after_the_link_check(tmp_path, monkeypatch):
    geocoder = FakeGeocoder()
    config = run_search(tmp_path, monkeypatch, geocoder)

    entries = load_locations(config.run_dir / SOURCE_LOCATIONS_FILE)
    assert geocoder.calls == [("Ballynagran, Coolbeg, Co. Wicklow", "IE")]
    assert entries["IRL|Ballynagran, Coolbeg, Co. Wicklow"]["basis"] == ROOFTOP
    assert entries["IRL|T 285 935"]["basis"] == "grid reference"

    # Arbitration reads the answers back, and never asks the geocoder itself.
    geocoder.calls.clear()
    run_arbitration(config)
    assert geocoder.calls == []
    bases = {row["stated_location"]: row["location_basis"] for row in csv.DictReader((config.run_dir / "evidence.csv").open())}
    assert bases == {ADDRESS["stated_location"]: ROOFTOP, GRID["stated_location"]: "grid reference"}

    run_config = {row[0]: row[1] for row in load_workbook(config.run_dir / "run_review.xlsx")["Run_Config"].iter_rows(values_only=True)}
    assert run_config["locations_stated"] == "2 (grid reference 1, rooftop 1)"
    assert run_config["locations_used"] == "2 (grid reference 1, rooftop 1)"
    assert run_config["geocoder"] == "fake"


def test_a_resumed_search_looks_up_only_what_is_new(tmp_path, monkeypatch):
    geocoder = FakeGeocoder()
    run_search(tmp_path, monkeypatch, geocoder)
    run_search(tmp_path, monkeypatch, geocoder)  # resumed: the cached response is reused
    assert len(geocoder.calls) == 1


def test_without_a_geocoder_addresses_stay_unlocated_and_run_config_says_so(tmp_path, monkeypatch):
    config = run_search(tmp_path, monkeypatch, None)
    entries = load_locations(config.run_dir / SOURCE_LOCATIONS_FILE)
    assert entries["IRL|Ballynagran, Coolbeg, Co. Wicklow"]["reason"] == NO_GEOCODER
    run_arbitration(config)
    run_config = {row[0]: row[1] for row in load_workbook(config.run_dir / "run_review.xlsx")["Run_Config"].iter_rows(values_only=True)}
    assert run_config["geocoder"] == "none configured"
    assert run_config["locations_used"] == "1 (grid reference 1)"


def test_no_locate_skips_the_step(tmp_path, monkeypatch):
    geocoder = FakeGeocoder()
    config = run_search(tmp_path, monkeypatch, geocoder, locate=False)
    assert not (config.run_dir / SOURCE_LOCATIONS_FILE).exists()
    assert geocoder.calls == []
    # Arbitration still reads coordinates and grid references itself; an address waits for locate.
    run_arbitration(config)
    reasons = {row["stated_location"]: row["location_reason"] for row in csv.DictReader((config.run_dir / "evidence.csv").open())}
    assert reasons == {ADDRESS["stated_location"]: NOT_LOCATED, GRID["stated_location"]: ""}
    run_config = {row[0]: row[1] for row in load_workbook(config.run_dir / "run_review.xlsx")["Run_Config"].iter_rows(values_only=True)}
    assert run_config["geocoder"] == "locate step not run"


def test_a_failed_lookup_is_reported_and_tried_again_next_time(tmp_path, monkeypatch, capsys):
    class Failing(FakeGeocoder):
        def geocode(self, address, iso2):
            super().geocode(address, iso2)
            raise ConnectionError("service unavailable")

    config = run_search(tmp_path, monkeypatch, Failing())
    assert "1 address lookup(s) failed and are not cached" in capsys.readouterr().out
    assert "IRL|Ballynagran, Coolbeg, Co. Wicklow" not in load_locations(config.run_dir / SOURCE_LOCATIONS_FILE)


def test_the_locate_command_backfills_an_older_run(tmp_path, monkeypatch):
    """An attribute-first run cached before the locate step existed: its found coordinates are the
    location the sources cited for both state."""
    from waste_ai_search.input_loader import write_csv_records
    from waste_ai_search.seed_source import seed_headers

    run_dir = tmp_path / "run"
    write_csv_records(
        run_dir / "seed.csv",
        [{"facility_id": "1", "site_id": "1", "internal_facility_id": "1", "year": "2022", "site_name": "Brady Road",
          "facility_name": "Brady Road", "country_iso3": "CAN", "iso3c_plus": "CAN", "latitude": "49.79", "longitude": "-97.16"}],
        seed_headers(),
    )
    raw = run_dir / "raw_foundry_responses"
    raw.mkdir(parents=True)
    cited = {**PERMIT, "url": "https://winnipeg.example/brady"}
    old = {"site_id": "1", "attributes": [
        {"attribute_name": name, "value": value, "unit": "decimal degrees", "value_basis": "Direct",
         "confidence_score": "High", "sources": [cited]}
        for name, value in (("found_latitude", "49.7958"), ("found_longitude", "-97.1601"))
    ]}
    (raw / "site_1.json").write_text(json.dumps(old))
    asked = []
    monkeypatch.setattr(search, "make_geocoder", lambda name: asked.append(name))

    assert cli.main(["locate", "--run-id", "run", "--run-dir", str(run_dir), "--geocoder", "nominatim"]) == 0
    assert asked == ["nominatim"]
    entries = load_locations(run_dir / SOURCE_LOCATIONS_FILE)
    assert entries["CAN|49.7958, -97.1601"]["basis"] == "coordinates"


def test_search_and_run_take_the_flags():
    args = cli.build_parser().parse_args(["search", "--no-locate", "--geocoder", "azure-maps"])
    config = cli.make_config(args)
    assert (config.locate, config.geocoder) == (False, "azure-maps")
    args = cli.build_parser().parse_args(["run"])
    config = cli.make_config(args)
    assert (config.locate, config.geocoder) == (True, "")
