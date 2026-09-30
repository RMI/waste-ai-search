"""WP-525: flag facilities that sources contradict as disposal sites.

A contradiction detector, not a verifier. The agent answers only when a source positively says a
site is something other than a waste facility, and the verdict only ever reaches a human.

The failure these tests exist to prevent is the one the ticket names: a closed, capped or
redeveloped landfill read as "not a landfill". Those sites keep emitting methane for decades and
are exactly what WasteMAP tracks, so "it's a park now" must never become a contradiction.

All offline: agent responses are hand-built payloads fed through the real extraction, resolution,
refresh, standardized-table and workbook code.
"""
from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from waste_ai_search.arbitrate import contradiction_rows, extract_evidence
from waste_ai_search.arbitration import needs_review, resolve
from waste_ai_search.prompt_builder import (
    build_site_prompt,
    needs_contradiction_check,
    requested_attributes,
)
from waste_ai_search.schema import (
    CONTRADICTION_HEADERS,
    GAP_FILL_ATTRIBUTES,
    NO_CONTRADICTION_FOUND,
    SITE_TYPE_CONTRADICTED,
    SITE_TYPE_CONTRADICTION,
    SITE_TYPE_CONTRADICTION_VALUES,
    STANDARDIZED_FACILITY_COLUMNS,
)


REGULATOR = {
    "source_title": "Mining concessions register",
    "publisher": "Provincial Ministry of Mines",
    "source_type": "Regulator",
    "publication_date": "2023-04-01",
    "url": "https://mines.example.gov/concession/4471",
    "quoted_evidence_short": "Granite quarry, active concession since 1998.",
}

WIKI = {
    "source_title": "Local places",
    "publisher": "Community wiki",
    "source_type": "Wiki or aggregator",
    "publication_date": "",
    "url": "https://wiki.example.org/place/4471",
    "quoted_evidence_short": "The site is a gravel pit.",
}

NEWS = {
    "source_title": "Old tip becomes a park",
    "publisher": "Regional Times",
    "source_type": "News report",
    "publication_date": "2019-06-01",
    "url": "https://news.example.com/tip-park",
    "quoted_evidence_short": "The former landfill was capped in 2009 and reopened as a park.",
}


def osm_site(**overrides):
    """An uncorroborated OSM facility: a name, a point, nothing else."""
    base = {
        "site_id": "4471",
        "internal_facility_id": "4471",
        "site_name": "Landfill",
        "country_iso3": "CAN",
        "latitude": "49.1",
        "longitude": "-122.3",
        "contributing_data_sources": "osm_2022",
        "reference_year": "2022",
    }
    base.update(overrides)
    return base


def attribute(name, value, sources, summary="summary", **extra):
    row = {
        "attribute_name": name,
        "value": value,
        "unit": "",
        "value_basis": "Direct",
        "confidence_score": "High",
        "value_date": "2023",
        "evidence_summary": summary,
        "sources": sources,
    }
    row.update(extra)
    return row


def arbitrate(site, *attributes):
    """Run one site's payload through extraction and resolution, as run_arbitration does."""
    payload = {"site_id": site["site_id"], "attributes": list(attributes)}
    by_attr, evidence, _sources, _warn = extract_evidence(site, payload, "run", "v", date(2026, 1, 1))
    names = list(requested_attributes(site)) + [n for n in by_attr if n not in requested_attributes(site)]
    resolved = [resolve(site, n, by_attr.get(n, []), reviewed_on=date(2026, 1, 1)) for n in names]
    return resolved, evidence


def row_for(resolved, name):
    return next(r for r in resolved if r["attribute_name"] == name)


# --- criterion 1: a positive contradiction ------------------------------------------------------
def test_a_source_describing_a_quarry_yields_contradicted_with_a_clickable_source():
    resolved, evidence = arbitrate(
        osm_site(),
        attribute(SITE_TYPE_CONTRADICTION, "Contradicted", [REGULATOR], summary="Active granite quarry."),
    )
    row = row_for(resolved, SITE_TYPE_CONTRADICTION)

    assert row["resolved_value"] == SITE_TYPE_CONTRADICTED
    assert row["winning_source_url"].startswith("https://")
    assert row["winning_source_tier"] == "Tier 1"


def test_even_a_tier_1_contradiction_is_never_auto_validated():
    """Without the override this is an ordinary Tier 1 empty-baseline fill, which auto-validates."""
    resolved, _ = arbitrate(osm_site(), attribute(SITE_TYPE_CONTRADICTION, "Contradicted", [REGULATOR]))
    row = row_for(resolved, SITE_TYPE_CONTRADICTION)

    assert row["validation_status"] == "Needs review"
    assert needs_review(row)
    assert "never signed off automatically" in row["resolution_rule"]


def test_the_override_is_what_prevents_auto_validation():
    """Control: the same Tier 1 source on an ordinary attribute does auto-validate."""
    resolved, _ = arbitrate(osm_site(), attribute("facility_type", "Dumpsite", [REGULATOR]))
    assert row_for(resolved, "facility_type")["validation_status"] == "Auto-validated"


# --- a low-tier contradiction -------------------------------------------------------------------
def test_a_low_tier_contradiction_still_reaches_review_marked_with_its_tier():
    """Tiered like any attribute: surfaced, but a reviewer can see it rests on a Tier 4 wiki."""
    resolved, _ = arbitrate(osm_site(), attribute(SITE_TYPE_CONTRADICTION, "Contradicted", [WIKI]))
    row = row_for(resolved, SITE_TYPE_CONTRADICTION)

    assert row["resolved_value"] == SITE_TYPE_CONTRADICTED
    assert row["winning_source_tier"] == "Tier 4"
    assert needs_review(row)


def test_a_higher_tier_contradiction_wins_over_a_lower_one():
    resolved, _ = arbitrate(
        osm_site(), attribute(SITE_TYPE_CONTRADICTION, "Contradicted", [WIKI, REGULATOR])
    )
    assert row_for(resolved, SITE_TYPE_CONTRADICTION)["winning_source_tier"] == "Tier 1"


# --- criterion 2: closure is not a contradiction ------------------------------------------------
def test_a_capped_landfill_reported_as_closed_is_no_contradiction():
    """The agent is told to report closure through facility_status and closing_year, and to omit
    the contradiction. That must resolve as no verdict, with closure captured as a normal fill."""
    resolved, evidence = arbitrate(
        osm_site(),
        attribute("facility_status", "Inactive", [NEWS]),
        attribute("closing_year", "2009", [NEWS]),
    )

    contradiction = row_for(resolved, SITE_TYPE_CONTRADICTION)
    assert contradiction["resolved_value"] == ""
    assert contradiction["resolution"] == "Not found"
    assert row_for(resolved, "facility_status")["resolved_value"] == "Inactive"
    assert row_for(resolved, "closing_year")["resolved_value"] == "2009"
    assert contradiction_rows(resolved, evidence) == []


def test_a_capped_landfill_wrongly_flagged_is_caught_by_the_closure_cross_check():
    """If the agent misreads "now a park" anyway, the reviewer is told to look twice."""
    resolved, evidence = arbitrate(
        osm_site(),
        attribute("facility_status", "Inactive", [NEWS]),
        attribute("closing_year", "2009", [NEWS]),
        attribute(SITE_TYPE_CONTRADICTION, "Contradicted", [NEWS], summary="Now a public park."),
    )
    [view] = contradiction_rows(resolved, evidence)

    assert view["closure_also_reported"].startswith("CHECK:")
    assert "facility_status = Inactive" in view["closure_also_reported"]
    assert "closing_year = 2009" in view["closure_also_reported"]


def test_the_guidance_tells_the_agent_closure_is_not_a_contradiction():
    prompt = build_site_prompt(osm_site(), attributes=requested_attributes(osm_site()))
    assert "IS STILL A WASTE" in prompt
    assert "was this ever a waste disposal site?" in prompt
    assert "facility_status = Inactive and closing_year" in prompt


# --- criterion 3: silence is never a negative ---------------------------------------------------
def test_no_contradicting_evidence_is_no_verdict_not_a_negative():
    resolved, evidence = arbitrate(osm_site())
    row = row_for(resolved, SITE_TYPE_CONTRADICTION)

    assert row["resolved_value"] == ""
    assert row["resolution"] == "Not found"
    assert contradiction_rows(resolved, evidence) == []


def test_a_site_the_agent_could_not_find_gets_no_verdict():
    """An empty attributes list - the agent's "I could not find this facility" - decides nothing."""
    resolved, evidence = arbitrate(osm_site())
    assert all(r["resolved_value"] == "" for r in resolved if r["attribute_name"] == SITE_TYPE_CONTRADICTION)


def test_an_explicit_no_contradiction_answer_is_dropped_rather_than_recorded():
    """"No contradiction found" carries no evidence of anything, so it maps to NULL."""
    resolved, evidence = arbitrate(
        osm_site(), attribute(SITE_TYPE_CONTRADICTION, "No contradiction found", [NEWS])
    )
    row = row_for(resolved, SITE_TYPE_CONTRADICTION)
    assert row["resolved_value"] == ""
    assert contradiction_rows(resolved, evidence) == []


def test_the_two_vocabulary_values_have_no_confirmed():
    """A "Confirmed" value would invite reasoning from absence."""
    assert SITE_TYPE_CONTRADICTION_VALUES == [SITE_TYPE_CONTRADICTED, NO_CONTRADICTION_FOUND]


# --- criterion 4: review only, nothing deleted or altered ---------------------------------------
def test_a_contradicted_verdict_never_reaches_the_standardized_table():
    from waste_ai_search.standardized import build_standard_record

    site = osm_site()
    resolved, _ = arbitrate(
        site,
        attribute("facility_type", "Dumpsite", [REGULATOR]),
        attribute(SITE_TYPE_CONTRADICTION, "Contradicted", [REGULATOR]),
    )
    record, _notes = build_standard_record(site, resolved, sha="deadbeef")

    assert SITE_TYPE_CONTRADICTION not in STANDARDIZED_FACILITY_COLUMNS
    assert SITE_TYPE_CONTRADICTION not in record
    # The facility's other findings are promoted exactly as they would be without the verdict.
    assert record["facility_type"] == "Dumpsite"


def test_an_unreviewed_verdict_is_never_merged_into_the_next_pass_seed():
    from waste_ai_search.seed_refresh import refresh_sites

    site = osm_site()
    resolved, _ = arbitrate(site, attribute(SITE_TYPE_CONTRADICTION, "Contradicted", [REGULATOR]))
    refreshed, _stats = refresh_sites([site], resolved, "run")

    assert SITE_TYPE_CONTRADICTION not in GAP_FILL_ATTRIBUTES
    assert not refreshed[0].get(SITE_TYPE_CONTRADICTION)


def test_the_contradictions_view_carries_source_tier_and_quote():
    resolved, evidence = arbitrate(
        osm_site(),
        attribute(SITE_TYPE_CONTRADICTION, "Contradicted", [REGULATOR], summary="Active granite quarry."),
    )
    [view] = contradiction_rows(resolved, evidence)

    assert set(CONTRADICTION_HEADERS) <= set(view)
    assert view["winning_source_tier"] == "Tier 1"
    assert view["evidence_summary"] == "Active granite quarry."
    assert "Granite quarry" in view["quoted_evidence_short"]
    assert view["closure_also_reported"] == ""
    assert view["validation_status"] == "Needs review"


def test_the_workbook_has_a_contradictions_tab(tmp_path):
    from waste_ai_search.workbook_io import write_review_workbook

    resolved, evidence = arbitrate(osm_site(), attribute(SITE_TYPE_CONTRADICTION, "Contradicted", [REGULATOR]))
    path = write_review_workbook(
        tmp_path / "r.xlsx",
        run_config=[],
        review_queue=[],
        leads=[],
        parse_warnings=[],
        contradictions=contradiction_rows(resolved, evidence),
    )
    ws = load_workbook(path)["Contradictions"]
    header = [c.value for c in ws[1]]
    assert header == CONTRADICTION_HEADERS
    assert ws.cell(2, header.index("verdict") + 1).value == SITE_TYPE_CONTRADICTED


# --- criterion 5: scope ---------------------------------------------------------------------------
def test_an_uncorroborated_osm_site_is_asked():
    assert needs_contradiction_check(osm_site())
    assert SITE_TYPE_CONTRADICTION in requested_attributes(osm_site())


def test_an_osm_site_with_any_corroborating_baseline_is_not_asked():
    """Once a pass has filled facility_type, the site has corroboration and is not asked again."""
    assert not needs_contradiction_check(osm_site(facility_type="Dumpsite"))
    assert not needs_contradiction_check(osm_site(annual_incoming_waste_metric_tonnes="12000"))


def test_area_from_the_same_osm_polygon_is_not_corroboration():
    """The area is measured from the very polygon in question, so it proves nothing."""
    assert needs_contradiction_check(osm_site(area_square_meters="101999.87"))


def test_a_site_with_a_second_source_is_not_asked():
    assert not needs_contradiction_check(osm_site(contributing_data_sources="osm_2022+eprtr_2022"))


def test_regulator_backed_and_gpw_sites_are_not_asked_as_written():
    """WP-525 scopes to OSM. gpw_2021 meets the same test and is a one-line decision, not a default."""
    assert not needs_contradiction_check(osm_site(contributing_data_sources="lmop_2024"))
    assert not needs_contradiction_check(osm_site(contributing_data_sources="gpw_2021"))


def test_a_corroborated_site_prompt_carries_none_of_the_guidance():
    """The other 12,644 facilities must not pay for a question they are never asked."""
    site = osm_site(contributing_data_sources="lmop_2024", facility_type="Sanitary Landfill")
    prompt = build_site_prompt(site, attributes=requested_attributes(site))
    assert SITE_TYPE_CONTRADICTION not in prompt
    assert "IS STILL A WASTE" not in prompt


# --- criterion 6: the premise -------------------------------------------------------------------
def test_the_prompt_no_longer_asserts_the_site_is_a_disposal_facility():
    prompt = build_site_prompt(osm_site(), attributes=requested_attributes(osm_site()))
    assert "for this waste disposal site" not in prompt
    assert "recorded as a waste disposal facility" in prompt
