"""Copilot review on WP-533: agent text must never become an Excel formula.

openpyxl stores a string starting with "=" as a formula, and Excel evaluates it on open. Every value
in the review workbook is agent or source text, so a quote such as '=HYPERLINK("http://evil","x")'
must be shown as text, not run.
"""
from __future__ import annotations

import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from waste_ai_search.workbook_io import write_sheet


@pytest.mark.parametrize("payload", [
    "=1+1",
    '=HYPERLINK("http://evil.example","click")',
    "=cmd|' /C calc'!A0",
    "=SUM(A1:A9)",
])
def test_text_starting_with_equals_is_stored_as_text(tmp_path, payload):
    wb = Workbook()
    write_sheet(wb.active, ["evidence_summary"], [{"evidence_summary": payload}])
    path = tmp_path / "w.xlsx"
    wb.save(path)

    cell = load_workbook(path).active["A2"]
    assert cell.value == payload
    assert cell.data_type == "s"
    with zipfile.ZipFile(path) as z:
        sheet = next(n for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
        assert "<f>" not in z.read(sheet).decode()


def test_ordinary_text_and_numbers_are_unchanged(tmp_path):
    wb = Workbook()
    write_sheet(wb.active, ["a", "b"], [{"a": "Active", "b": 42}])
    path = tmp_path / "w.xlsx"
    wb.save(path)
    ws = load_workbook(path).active
    assert ws["A2"].value == "Active" and ws["B2"].value == 42
