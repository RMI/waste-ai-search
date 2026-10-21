"""A found location is one value, read from where the sources say the facility is.

Since build step 4 of the source-matching plan, every source states the facility's location as it
writes it, and the locate step turns that into a point (locate.py). The site's found_location is
resolved from its sources' usable points - coordinates, a grid reference or a rooftop address - by
the usual rules: the most credible source, then the newest. It is no longer two coordinates the
agent returned and arbitration joined, so the tests of that join (combine_location, and half a
location listed as a lead) are gone on purpose. A cached attribute-first response still arbitrates:
its found coordinates become the location stated by the sources cited for both.

A location that differs from ours, by however little, goes to review like any other value.

Changed on purpose for build step 5: a source is filed under the site it is about by where it
says the facility is (match.py), so El Hongo now has a location of ours, and a source more than
2 km from it is a new-facility candidate, not one of its values. Brazil's municipality-centre site
takes a location up to 25 km away only where the source's name matches it.
"""
from __future__ import annotations

import csv
import json

from openpyxl import load_workbook

from waste_ai_search.arbitrate import run_arbitration
from waste_ai_search.input_loader import write_csv_records
from waste_ai_search.run_context import PipelineConfig
from waste_ai_search.schema import LOCATION_ATTRIBUTE, REVIEW_QUEUE_HEADERS
from waste_ai_search.seed_refresh import refresh_sites
from waste_ai_search.workbook_io import REVIEW_INSTRUCTIONS


def test_refresh_seed_merges_a_found_location_into_both_coordinates():
    site = {"site_id": "21", "latitude": "32.0", "longitude": "-116.0", "ai_filled_fields": ""}
    location = {
        "site_id": "21",
        "attribute_name": LOCATION_ATTRIBUTE,
        "resolved_value": "32.51, -116.21",
        "resolution": "Overrides baseline",
        "winning_source_tier": "Tier 2",
    }

    (refreshed,), stats = refresh_sites([site], [location], "run", merge_identity=True)
    assert (refreshed["latitude"], refreshed["longitude"]) == ("32.51", "-116.21")
    # Stamped per coordinate, the way the next pass reads them.
    assert refreshed["ai_filled_fields"] == "found_latitude@Tier 2; found_longitude@Tier 2"
    assert stats["values_merged"] == 2

    (unchanged,), _ = refresh_sites([site], [location], "run")
    assert (unchanged["latitude"], unchanged["longitude"]) == ("32.0", "-116.0")


def test_a_correction_can_carry_its_own_date():
    assert REVIEW_QUEUE_HEADERS.index("corrected_value_date") == REVIEW_QUEUE_HEADERS.index("corrected_value") + 1
    queue = next(row for row in REVIEW_INSTRUCTIONS if row["tab_name"] == "Review_Queue")
    assert "corrected_value_date" in queue["editable"]


# --- end to end -----------------------------------------------------------------------------
NEWS = {
    "source_title": "El Hongo dump to close",
    "publisher": "Diario local",
    "source_type": "News report",
    "publication_date": "2023-05-01",
    "url": "https://news.example.mx/el-hongo",
}
REGULATOR = {**NEWS, "publisher": "SEMARNAT", "source_type": "Regulator", "url": "https://gob.example.mx/el-hongo"}


def arbitrate_el_hongo(tmp_path, response, **seed):
    """Arbitrate one site, at OSM's point unless given another, against one cached response."""
    from waste_ai_search.seed_source import seed_headers

    run_dir = tmp_path / "run"
    site = {
        "facility_id": "21", "site_id": "21", "internal_facility_id": "21", "year": "2016",
        "site_name": "EL HONGO", "facility_name": "EL HONGO", "country_iso3": "MEX", "iso3c_plus": "MEX",
        "contributing_data_sources": "osm_2022", "reference_year": "2016",
        "latitude": "32.5163", "longitude": "-116.2141", **seed,
    }
    write_csv_records(run_dir / "seed.csv", [site], seed_headers())
    raw = run_dir / "raw_foundry_responses"
    raw.mkdir(parents=True)
    (raw / "site_21.json").write_text(json.dumps({"site_id": "21", **response}))

    run_arbitration(PipelineConfig(input_csv=None, run_dir=run_dir, run_id="run", dataset_version="v"))
    return run_dir


def old_response(coordinates, source=NEWS):
    """An attribute-first response, as batch 1 cached them."""
    return {"attributes": [
        {"attribute_name": name, "value": value, "unit": "decimal degrees", "value_basis": "Direct",
         "confidence_score": "High", "value_date": "2023", "evidence_summary": "The article gives the site's coordinates.",
         "sources": [source]}
        for name, value in coordinates
    ]}


def stated(*sources):
    """A source-first response: each (source, stated location, kind), with no values."""
    return {"sources": [
        {**source, "stated_location": location, "stated_location_kind": kind, "values": []}
        for source, location, kind in sources
    ]}


def rows(run_dir, name):
    return list(csv.DictReader((run_dir / name).open(encoding="utf-8-sig")))


def test_an_older_response_still_gives_one_location_row(tmp_path):
    run_dir = arbitrate_el_hongo(tmp_path, old_response([("found_latitude", "32.5122"), ("found_longitude", "-116.2134")]))

    # A Tier 3 source disagrees with OSM's point, so the point waits for review as one row.
    queue = rows(run_dir, "review_queue.csv")
    assert [(row["attribute_name"], row["resolved_value"], row["resolution"]) for row in queue] == [
        (LOCATION_ATTRIBUTE, "32.5122, -116.2134", "Conflict - needs review")
    ]
    [evidence] = rows(run_dir, "evidence.csv")
    assert (evidence["stated_location"], evidence["location_basis"]) == ("32.5122, -116.2134", "coordinates")

    ws = load_workbook(run_dir / "run_review.xlsx")["Review_Queue"]
    header = [cell.value for cell in ws[1]]
    assert header.index("corrected_value_date") == header.index("corrected_value") + 1


def test_a_regulators_location_a_few_hundred_metres_from_ours_is_reviewed(tmp_path):
    # The permit's point is about 0.25 km from OSM's. The 1 km rule once settled it as the same
    # place, auto-validated; now it is a value like any other, and the SME decides.
    run_dir = arbitrate_el_hongo(
        tmp_path, stated((REGULATOR, "32.5122, -116.2000", "coordinates")), latitude="32.51", longitude="-116.2",
    )
    assert [(row["attribute_name"], row["resolved_value"], row["resolution"]) for row in rows(run_dir, "review_queue.csv")] == [
        (LOCATION_ATTRIBUTE, "32.5122, -116.2", "Overrides baseline")
    ]


def test_the_most_credible_source_wins_then_the_newest(tmp_path):
    older_permit = {**REGULATOR, "publication_date": "2015-01-01", "url": "https://gob.example.mx/2015"}
    newer_permit = {**REGULATOR, "publication_date": "2021-01-01", "url": "https://gob.example.mx/2021"}
    newest_news = {**NEWS, "publication_date": "2024-06-01"}
    run_dir = arbitrate_el_hongo(tmp_path, stated(
        (older_permit, "32.5100, -116.2100", "coordinates"),
        (newest_news, "32.6000, -116.3000", "coordinates"),
        (newer_permit, "32.5200, -116.2200", "coordinates"),
    ))
    [location] = rows(run_dir, "review_queue.csv")
    assert (location["resolved_value"], location["winning_source_url"]) == ("32.52, -116.22", newer_permit["url"])
    assert location["value_date"] == "2021-01-01"
    # The news story places it 12 km away, where none of our sites is: a new-facility candidate.
    [candidate] = rows(run_dir, "new_facility_candidates.csv")
    assert (candidate["source_urls"], candidate["nearest_site_id"]) == (newest_news["url"], "21")


def test_coarse_coordinates_are_not_a_location_value_nor_a_lead(tmp_path):
    run_dir = arbitrate_el_hongo(tmp_path, old_response([("found_latitude", "32.51"), ("found_longitude", "-116.21")]))
    assert rows(run_dir, "review_queue.csv") == []
    assert not [lead for lead in rows(run_dir, "supplementary_leads.csv") if lead["attribute_name"] == LOCATION_ATTRIBUTE]
    [evidence] = rows(run_dir, "evidence.csv")
    assert (evidence["location_basis"], evidence["promotion_eligible"]) == ("hint", "FALSE")
    assert evidence["exclusion_reason"].startswith("A hint, not a location value")


def test_half_a_location_is_no_location_and_no_lead(tmp_path):
    """Reversed on purpose: a lone coordinate used to be listed under Leads as half a location."""
    wiki = {**NEWS, "publisher": "Wikimapia", "source_type": "Wiki or aggregator", "url": "https://wiki.example/el-hongo"}
    run_dir = arbitrate_el_hongo(tmp_path, old_response([("found_latitude", "32.5122")], source=wiki))
    assert rows(run_dir, "review_queue.csv") == []
    assert rows(run_dir, "supplementary_leads.csv") == []
    assert rows(run_dir, "evidence.csv") == []


def test_a_grid_reference_becomes_the_found_location(tmp_path):
    from waste_ai_search.seed_source import seed_headers

    run_dir = tmp_path / "run"
    site = {"facility_id": "5", "site_id": "5", "internal_facility_id": "5", "year": "2022", "site_name": "Ben Nevis Tip",
            "facility_name": "Ben Nevis Tip", "country_iso3": "GBR", "iso3c_plus": "GBR",
            "contributing_data_sources": "osm_2022", "reference_year": "2022", "latitude": "56.79", "longitude": "-5.01"}
    write_csv_records(run_dir / "seed.csv", [site], seed_headers())
    raw = run_dir / "raw_foundry_responses"
    raw.mkdir(parents=True)
    sepa = {**REGULATOR, "publisher": "SEPA", "url": "https://sepa.example/permit"}
    (raw / "site_5.json").write_text(json.dumps({"site_id": "5", **stated((sepa, "NN 16667 71283", "grid reference"))}))
    run_arbitration(PipelineConfig(input_csv=None, run_dir=run_dir, run_id="run", dataset_version="v"))

    [location] = rows(run_dir, "review_queue.csv")
    assert location["attribute_name"] == LOCATION_ATTRIBUTE
    lat, lon = (float(part) for part in location["resolved_value"].split(","))
    assert abs(lat - 56.79687) < 0.0005 and abs(lon + 5.00367) < 0.0005
    [evidence] = rows(run_dir, "evidence.csv")
    assert evidence["location_basis"] == "grid reference"


def test_where_our_coordinates_are_trusted_source_locations_are_recorded_not_queued(tmp_path):
    run_dir = arbitrate_el_hongo(
        tmp_path, stated((REGULATOR, "32.5122, -116.2000", "coordinates")), latitude="32.51", longitude="-116.2",
        attribute_sources="found_latitude@mexico_inegi_2016; found_longitude@mexico_inegi_2016",
    )
    assert rows(run_dir, "review_queue.csv") == []
    assert rows(run_dir, "supplementary_leads.csv") == []
    [evidence] = rows(run_dir, "evidence.csv")
    assert (evidence["attribute_name"], evidence["location_basis"], evidence["promotion_eligible"]) == (
        LOCATION_ATTRIBUTE, "coordinates", "FALSE",
    )
    assert evidence["exclusion_reason"].startswith("Recorded only")


def test_a_source_with_only_a_location_is_a_source(tmp_path):
    run_dir = arbitrate_el_hongo(tmp_path, stated((REGULATOR, "32.5122, -116.2000", "coordinates")))
    [source] = rows(run_dir, "sources.csv")
    [evidence] = rows(run_dir, "evidence.csv")
    assert source["source_id"] == evidence["source_id"]


def test_a_value_row_carries_its_sources_facility_and_location(tmp_path):
    source = {**REGULATOR, "facility_name_in_source": "Relleno Sanitario El Hongo",
              "facility_name_english": "El Hongo Sanitary Landfill",
              "stated_location": "32.5122, -116.2000", "stated_location_kind": "coordinates",
              "values": [{"attribute_name": "operator", "value": "Ayuntamiento de Tecate", "unit": "",
                          "value_basis": "Direct", "confidence_score": "High", "value_date": "2023",
                          "evidence_summary": "The permit names the operator."}]}
    run_dir = arbitrate_el_hongo(tmp_path, {"sources": [source]})
    operator = next(row for row in rows(run_dir, "evidence.csv") if row["attribute_name"] == "operator")
    assert operator["facility_name_in_source"] == "Relleno Sanitario El Hongo"
    assert operator["facility_name_english"] == "El Hongo Sanitary Landfill"
    assert (operator["location_basis"], operator["located_latitude"], operator["located_longitude"]) == (
        "coordinates", "32.5122", "-116.2",
    )


def test_a_coordinates_only_site_gets_its_found_location_from_a_source(tmp_path):
    """Brazil's inexact SINIR sites are searched for sources that say where the dump is. The permit
    places it 8 km from the municipality centre we hold, beyond the 2 km match but within the 25 km
    an inexact location allows a source whose name matches."""
    from waste_ai_search.seed_source import seed_headers

    run_dir = tmp_path / "run"
    name = "Lixão de Porto Nacional"
    site = {"facility_id": "9", "site_id": "9", "internal_facility_id": "9", "year": "2024", "site_name": name,
            "facility_name": name, "country_iso3": "BRA", "iso3c_plus": "BRA", "is_location_exact": "FALSE",
            "contributing_data_sources": "sinir_2024", "reference_year": "2024",
            "latitude": "-10.1800", "longitude": "-48.3300"}
    write_csv_records(run_dir / "seed.csv", [site], seed_headers())
    raw = run_dir / "raw_foundry_responses"
    raw.mkdir(parents=True)
    permit = {**REGULATOR, "publisher": "NATURATINS", "url": "https://naturatins.example.br/licenca/123",
              "facility_name_in_source": "Lixão Municipal de Porto Nacional"}
    (raw / "site_9.json").write_text(json.dumps({"site_id": "9", **stated((permit, "-10.2481, -48.3612", "coordinates"))}))
    run_arbitration(PipelineConfig(input_csv=None, run_dir=run_dir, run_id="run", dataset_version="v"))

    [location] = rows(run_dir, "review_queue.csv")
    assert (location["attribute_name"], location["resolved_value"]) == (LOCATION_ATTRIBUTE, "-10.2481, -48.3612")
    assert location["match"].startswith("Name nearby:") and "coordinate correction" in location["match"]
    assert location["location_basis"] == "coordinates"
