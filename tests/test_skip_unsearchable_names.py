"""Sites whose name gives a search nothing to look for are not searched (step 3 of
docs/plans/2026-10-15-skip-unsearchable-names.md). All four classes are skipped, at site selection,
beside the US and Brazil rule and in the same way: the run says what it dropped, an explicit request
wins, and the workbook's Run_Config records it."""
from __future__ import annotations

import contextlib
import csv
import io
from dataclasses import replace

import pytest

from waste_ai_search.run_context import PipelineConfig
from waste_ai_search.searchable_names import (
    GENERIC,
    MINING,
    NAMED,
    NO_NAME,
    NOT_A_NAME,
    SKIPPED_SITES_FILE,
    describe_counts,
    name_class,
)

# (site_id, site_name, original_site_name): one of each class, then a site named after a mining
# word that only its untranslated name shows.
NAMES = [
    ("1", "Hartland Landfill", ""),
    ("2", "OpenStreetMap Landfill", ""),
    ("3", "", ""),
    ("4", "update the tag landuse after survey", ""),
    ("5", "domestic landfill", ""),
    ("6", "Overburden Blade", ""),
    ("7", "Northern dump", "Отвал Северный"),
]


def write_seed(tmp_path, names=NAMES, country="CAN"):
    path = tmp_path / "seed.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["site_id", "site_name", "original_site_name", "country_iso3", "latitude", "longitude"])
        for site_id, name, original in names:
            writer.writerow([site_id, name, original, country, "49.1", "-123.2"])
    return path


def select(seed, **kwargs):
    from waste_ai_search.input_loader import load_sites
    from waste_ai_search.search import select_sites

    sites, _headers = load_sites(seed)
    printed = io.StringIO()
    with contextlib.redirect_stdout(printed):
        chosen = select_sites(sites, PipelineConfig(input_csv=seed, run_dir=seed.parent, run_id="t", **kwargs))
    return [s["site_id"] for s in chosen], printed.getvalue()


# --- selection --------------------------------------------------------------------------------
def test_all_four_classes_are_skipped_by_default_and_named_sites_are_searched(tmp_path):
    ids, _printed = select(write_seed(tmp_path))
    assert ids == ["1"]


def test_the_run_says_how_many_it_skipped_in_each_class(tmp_path):
    _ids, printed = select(write_seed(tmp_path))
    assert (
        "Skipped 6 site(s) whose name gives a search nothing to look for (no name 2, not a name 1, "
        "generic label 1, mining feature 2; pass --include-unsearchable-names to search them)."
    ) in printed


def test_nothing_is_said_when_nothing_is_skipped(tmp_path):
    _ids, printed = select(write_seed(tmp_path, names=[("1", "Hartland Landfill", "")]))
    assert "whose name gives a search nothing to look for" not in printed


def test_listing_site_ids_searches_them_anyway(tmp_path):
    """Naming a site is an explicit request, as with the country rule."""
    ids, printed = select(write_seed(tmp_path), site_ids=["2", "5"])
    assert ids == ["2", "5"]
    assert "Skipped" not in printed


def test_the_flag_searches_them_anyway(tmp_path):
    ids, _printed = select(write_seed(tmp_path), include_unsearchable_names=True)
    assert ids == ["1", "2", "3", "4", "5", "6", "7"]


def test_naming_the_country_does_not_name_its_unnamed_sites(tmp_path):
    """--iso3 lifts the country rule, not this one: it names a country, not these sites."""
    ids, _printed = select(write_seed(tmp_path, country="USA"), iso3=["USA"])
    assert ids == ["1"]


def test_a_batch_counts_what_it_will_search_after_the_skip(tmp_path):
    _ids, printed = select(write_seed(tmp_path), batch=1)
    assert "Batch 1: 1 site(s) in 1 country(ies)." in printed


def test_a_pilot_sample_includes_them_only_with_the_flag(tmp_path):
    """--pilot-size deliberately draws generic-named hard cases, but only when they are searched."""
    names = [(str(i), f"Hillside Landfill {i}", "") for i in range(1, 6)] + [("6", "City Landfill", "")]
    seed = write_seed(tmp_path, names=names)

    default, _printed = select(seed, pilot_size=6)
    assert "6" not in default

    from waste_ai_search.input_loader import load_sites
    from waste_ai_search.search import select_sites

    sites, _headers = load_sites(seed)
    with contextlib.redirect_stdout(io.StringIO()):
        opted_in = select_sites(
            sites, PipelineConfig(input_csv=seed, run_dir=tmp_path, run_id="t", pilot_size=6, include_unsearchable_names=True)
        )
    reasons = {s["site_id"]: s["pilot_selection_reason"] for s in opted_in}
    assert reasons["6"] == "generic_name_hard_case"


def test_the_cli_flag_reaches_the_config_for_search_and_run():
    from waste_ai_search.cli import build_parser, make_config

    for command in ("search", "run"):
        args = build_parser().parse_args([command, "--run-id", "t", "--include-unsearchable-names"])
        assert make_config(args).include_unsearchable_names
        args = build_parser().parse_args([command, "--run-id", "t"])
        assert not make_config(args).include_unsearchable_names


def test_the_measurement_still_sees_the_sites_the_rule_skips():
    """The forecast and --list-sites measure these classes, so they select without the rule."""
    from waste_ai_search import name_yield as ny

    sites = [
        {"site_id": "1", "site_name": "OpenStreetMap Landfill", "country_iso3": "CAN"},
        {"site_id": "2", "site_name": "Hartland Landfill", "country_iso3": "CAN"},
    ]
    with contextlib.redirect_stdout(io.StringIO()):
        selections = ny.batch_selections(sites)
    assert [s["site_id"] for s in selections[1]] == ["1", "2"]


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        # Initials and route codes say which site it is (local 15 Oct scan: AUS 10836, USA 3525,
        # 5141, 5846, 1574).
        ("K & W Landfill", NAMED),
        ("R & B Landfill", NAMED),
        ("Sr 85 Landfill", NAMED),
        ("Gb Landfill", NAMED),
        ("Bj Landfill", NAMED),
        # Numbers, lone letters, roman numerals and "no" only number a site or cell.
        ("Dump 1", GENERIC),
        ("Landfill No 2", GENERIC),
        ("Landfill A", GENERIC),
        ("Landfill II", GENERIC),
        ("Garbage dump", GENERIC),
    ],
)
def test_initials_and_route_codes_count_as_a_name(name, expected):
    assert name_class(name) == expected


@pytest.mark.parametrize(
    ("name", "original", "expected"),
    [
        # The translation lost a name the source recorded; the agent searches the original first.
        ("", "Свалка у деревни Ивановка", NAMED),
        ("OpenStreetMap Landfill", "Hartland", NAMED),
        ("OpenStreetMap Landfill", "", NO_NAME),
        ("", "Отвал Северный", NO_NAME),  # the original is a spoil heap, so still skipped
        # Known limitation: the generic check knows only English words, so it can't tell a
        # distinctive original from a generic one such as Полигон ТБО, and reads the English name.
        ("Solid Waste Landfill", "Полигон ТБО «Новосёлки»", GENERIC),
    ],
)
def test_an_empty_english_name_with_a_source_spelling_is_named(name, original, expected):
    assert name_class(name, original) == expected


@pytest.mark.parametrize(
    ("classes", "expected"),
    [
        ([], "0"),
        ([GENERIC, NO_NAME, GENERIC], "3 (no name 1, generic label 2)"),
        ([MINING, NOT_A_NAME], "2 (not a name 1, mining feature 1)"),
    ],
)
def test_counts_are_described_in_class_order(classes, expected):
    assert describe_counts(classes) == expected


# --- the skipped list and Run_Config ----------------------------------------------------------
class _Quiet:
    def __init__(self):
        self.searched = []

    def search_site(self, site, prompt):  # noqa: ARG002
        self.searched.append(site["site_id"])
        return {"attributes": [], "search_notes": "Searched thoroughly and found no source-backed data."}


def search(config, monkeypatch):
    from waste_ai_search import search as se

    client = _Quiet()
    monkeypatch.setattr(se, "get_client", lambda config: client)
    with contextlib.redirect_stdout(io.StringIO()):
        se.run_search(config)
    return client.searched


def skipped_rows(run_dir):
    with (run_dir / SKIPPED_SITES_FILE).open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def run_config(config):
    from openpyxl import load_workbook

    from waste_ai_search.arbitrate import run_arbitration

    with contextlib.redirect_stdout(io.StringIO()):
        run_arbitration(config)
    sheet = load_workbook(config.run_dir / f"{config.run_id}_review.xlsx")["Run_Config"]
    return {row[0]: row[1] for row in sheet.iter_rows(values_only=True)}


def base_config(tmp_path, **kwargs):
    return PipelineConfig(
        input_csv=write_seed(tmp_path), run_dir=tmp_path / "run", run_id="run",
        site_delay_seconds=0, check_links=False, **kwargs,
    )


def test_the_skipped_sites_are_listed_in_the_run_folder_with_their_class(tmp_path, monkeypatch):
    config = base_config(tmp_path)
    assert search(config, monkeypatch) == ["1"]

    rows = skipped_rows(config.run_dir)
    assert [(row["site_id"], row["name_class"]) for row in rows] == [
        ("2", NO_NAME), ("3", NO_NAME), ("4", NOT_A_NAME), ("5", GENERIC), ("6", MINING), ("7", MINING),
    ]
    assert rows[0]["reason"] == "ingestion's filler for an unnamed OSM polygon"
    assert rows[5]["original_site_name"] == "Отвал Северный"
    assert rows[0]["country_iso3"] == "CAN"


def test_run_config_records_what_was_skipped(tmp_path, monkeypatch):
    config = base_config(tmp_path)
    search(config, monkeypatch)

    settings = run_config(config)
    assert settings["unsearchable_names_skipped"] == "6 (no name 2, not a name 1, generic label 1, mining feature 2)"
    assert settings["unsearchable_names_searched"] == "0"


def test_a_search_of_listed_ids_keeps_the_runs_list_but_takes_its_sites_off(tmp_path, monkeypatch):
    """A retry or the gas follow-up names its sites; it must not wipe the run's own list. A site it
    searches is no longer skipped, so it is counted once, as searched."""
    config = base_config(tmp_path)
    search(config, monkeypatch)
    assert search(replace(config, site_ids=["2"]), monkeypatch) == ["2"]

    assert [row["site_id"] for row in skipped_rows(config.run_dir)] == ["3", "4", "5", "6", "7"]
    settings = run_config(config)
    assert settings["unsearchable_names_skipped"] == "5 (no name 1, not a name 1, generic label 1, mining feature 2)"
    assert settings["unsearchable_names_searched"] == "1 (no name 1)"


def test_a_listed_search_with_no_list_yet_starts_none(tmp_path, monkeypatch):
    config = base_config(tmp_path)
    search(replace(config, site_ids=["1"]), monkeypatch)

    assert not (config.run_dir / SKIPPED_SITES_FILE).exists()


def test_with_the_flag_nothing_is_skipped_and_run_config_says_what_was_searched(tmp_path, monkeypatch):
    config = base_config(tmp_path)
    search(config, monkeypatch)  # an earlier attempt with the rule on
    assert len(search(replace(config, include_unsearchable_names=True), monkeypatch)) == 6  # site 1 is cached

    assert skipped_rows(config.run_dir) == []
    settings = run_config(config)
    assert settings["unsearchable_names_skipped"] == "0"
    assert settings["unsearchable_names_searched"] == "6 (no name 2, not a name 1, generic label 1, mining feature 2)"


# Two countries, each with one named site and two the rule skips.
TWO_COUNTRIES = [
    ("1", "Hartland Landfill", "CAN"),
    ("2", "OpenStreetMap Landfill", "CAN"),
    ("3", "domestic landfill", "CAN"),
    ("4", "Olusosun Dumpsite", "NGA"),
    ("5", "Dump", "NGA"),
    ("6", "OpenStreetMap Landfill", "NGA"),
]


def two_country_config(tmp_path):
    path = tmp_path / "seed.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["site_id", "site_name", "country_iso3", "latitude", "longitude"])
        for site_id, name, country in TWO_COUNTRIES:
            writer.writerow([site_id, name, country, "1.0", "2.0"])
    return PipelineConfig(
        input_csv=path, run_dir=tmp_path / "run", run_id="run", site_delay_seconds=0, check_links=False
    )


def test_a_narrower_search_under_the_same_run_id_adds_to_the_list(tmp_path, monkeypatch):
    """Resuming one country must not drop the other country's skipped sites from the record."""
    config = two_country_config(tmp_path)
    search(config, monkeypatch)
    search(replace(config, iso3=["CAN"]), monkeypatch)

    assert [row["site_id"] for row in skipped_rows(config.run_dir)] == ["2", "3", "5", "6"]
    assert run_config(config)["unsearchable_names_skipped"] == "4 (no name 2, generic label 2)"


def test_the_flag_on_one_country_leaves_the_others_skips_on_record(tmp_path, monkeypatch):
    config = two_country_config(tmp_path)
    search(config, monkeypatch)
    assert sorted(search(replace(config, iso3=["NGA"], include_unsearchable_names=True), monkeypatch)) == ["5", "6"]

    assert [row["site_id"] for row in skipped_rows(config.run_dir)] == ["2", "3"]
    settings = run_config(config)
    assert settings["unsearchable_names_skipped"] == "2 (no name 1, generic label 1)"
    assert settings["unsearchable_names_searched"] == "2 (no name 1, generic label 1)"


def test_a_resume_after_the_flag_does_not_list_sites_already_searched(tmp_path, monkeypatch):
    """The resume would skip them by name, but the run has already searched them."""
    config = two_country_config(tmp_path)
    search(replace(config, include_unsearchable_names=True), monkeypatch)
    assert search(config, monkeypatch) == []  # everything is cached

    assert skipped_rows(config.run_dir) == []
    settings = run_config(config)
    assert settings["unsearchable_names_skipped"] == "0"
    assert settings["unsearchable_names_searched"] == "4 (no name 2, generic label 2)"


def test_a_limit_counts_every_site_the_rule_dropped_from_the_selection(tmp_path, monkeypatch):
    """The rule runs before --limit, so the limit takes searchable sites; the list covers the
    whole selection, not only the sites the limit would have reached."""
    config = two_country_config(tmp_path)
    assert search(replace(config, limit=1), monkeypatch) == ["1"]

    assert len(skipped_rows(config.run_dir)) == 4


def test_a_run_from_before_the_rule_says_it_was_not_recorded(tmp_path, monkeypatch):
    config = base_config(tmp_path)
    search(config, monkeypatch)
    (config.run_dir / SKIPPED_SITES_FILE).unlink()

    assert run_config(config)["unsearchable_names_skipped"].startswith("not recorded")
