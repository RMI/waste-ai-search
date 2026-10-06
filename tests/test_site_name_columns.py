"""WP-531 item 1: every SME-facing output shows the translated and original site names together.

So a reviewer can check results against the source's own spelling and record translation problems
in `translation_note`, the only new editable column.

The end-to-end test builds a run directory - a pinned seed and one cached agent response - and runs
the real arbitration over it. It is offline: the seed snapshot means no database is touched.
"""
from __future__ import annotations

import csv
import json

import pytest
from openpyxl import load_workbook

from waste_ai_search.arbitrate import attach_site_names, run_arbitration
from waste_ai_search.input_loader import write_csv_records
from waste_ai_search.run_context import PipelineConfig
from waste_ai_search.schema import (
    EVIDENCE_HEADERS,
    RESOLVED_HEADERS,
    REVIEW_QUEUE_HEADERS,
    SITE_NAME_COLUMNS,
    SUPPLEMENTARY_LEADS_HEADERS,
)


ALL_OUTPUTS = {
    "evidence.csv": EVIDENCE_HEADERS,
    "resolved.csv": RESOLVED_HEADERS,
    "Review_Queue": REVIEW_QUEUE_HEADERS,
    "Leads": SUPPLEMENTARY_LEADS_HEADERS,
}


# --- headers --------------------------------------------------------------------------------
@pytest.mark.parametrize("output,headers", ALL_OUTPUTS.items())
def test_every_output_carries_the_three_name_columns_side_by_side(output, headers):
    start = headers.index("site_name")
    assert headers[start:start + 3] == SITE_NAME_COLUMNS, output


def test_translation_note_sits_beside_the_names_in_the_review_queue():
    start = REVIEW_QUEUE_HEADERS.index("site_name")
    assert REVIEW_QUEUE_HEADERS[start + 3] == "translation_note"


def test_translation_note_is_only_in_the_review_queue():
    """An editable column belongs where the decision is recorded, not in the read-only outputs."""
    for output, headers in ALL_OUTPUTS.items():
        if output != "Review_Queue":
            assert "translation_note" not in headers, output


# --- filling the values ---------------------------------------------------------------------
SITES = {
    "21": {"site_id": "21", "site_name": "THE FUNGUS", "original_site_name": "EL HONGO", "source_language": "es"},
    "7": {"site_id": "7", "site_name": "Greenview Landfill", "original_site_name": "", "source_language": ""},
}


def test_names_are_filled_from_the_seed_by_site():
    rows = [{"site_id": "21"}, {"site_id": "7"}]
    attach_site_names(rows, SITES)
    assert rows[0] == {"site_id": "21", "site_name": "THE FUNGUS", "original_site_name": "EL HONGO", "source_language": "es"}
    assert rows[1]["original_site_name"] == ""  # untranslated source: blank, not invented


def test_an_existing_site_name_is_not_overwritten():
    rows = [{"site_id": "21", "site_name": "Already set"}]
    attach_site_names(rows, SITES)
    assert rows[0]["site_name"] == "Already set"
    assert rows[0]["original_site_name"] == "EL HONGO"


def test_a_row_for_an_unknown_site_gets_blanks_rather_than_failing():
    rows = [{"site_id": "999"}]
    attach_site_names(rows, SITES)
    assert rows[0]["original_site_name"] == "" and rows[0]["source_language"] == ""


# --- end to end -----------------------------------------------------------------------------
def test_a_real_arbitration_writes_both_names_everywhere(tmp_path):
    from waste_ai_search.seed_source import seed_headers

    run_dir = tmp_path / "run"
    site = {
        "facility_id": "21", "site_id": "21", "internal_facility_id": "21", "year": "2016",
        "site_name": "THE FUNGUS", "facility_name": "THE FUNGUS",
        "original_site_name": "EL HONGO", "source_language": "es", "name_data_source": "mexico_inegi_2016",
        "country_iso3": "MEX", "iso3c_plus": "MEX", "latitude": "32.5", "longitude": "-116.2",
        "contributing_data_sources": "mexico_inegi_2016", "reference_year": "2016",
    }
    write_csv_records(run_dir / "seed.csv", [site], seed_headers())

    response = {
        "site_id": "21",
        "attributes": [{
            "attribute_name": "facility_status",
            "value": "Inactive",
            "value_basis": "Direct",
            "confidence_score": "High",
            "value_date": "2023",
            "evidence_summary": "Clandestine dump closed by the municipality.",
            "sources": [{
                "source_title": "Cierran basurero clandestino de El Hongo",
                "publisher": "Diario local",
                "source_type": "News report",
                "publication_date": "2023-05-01",
                "url": "https://news.example.mx/el-hongo",
            }],
        }],
        "unverified_leads": [],
    }
    raw = run_dir / "raw_foundry_responses"
    raw.mkdir(parents=True)
    (raw / "site_21.json").write_text(json.dumps(response))

    config = PipelineConfig(input_csv=None, run_dir=run_dir, run_id="run", dataset_version="v")
    run_arbitration(config)

    # The value is Tier 3, so it waits for review: resolved.csv keeps only auto-validated rows.
    for name in ("review_queue.csv", "evidence.csv"):
        row = next(csv.DictReader((run_dir / name).open()))
        assert (row["site_name"], row["original_site_name"], row["source_language"]) == ("THE FUNGUS", "EL HONGO", "es"), name
    # A real arbitration fills the seed location for pasting into Google Maps.
    assert next(csv.DictReader((run_dir / "review_queue.csv").open()))["baseline_coordinates"] == "32.500000, -116.200000"

    ws = load_workbook(run_dir / "run_review.xlsx")["Review_Queue"]
    header = [c.value for c in ws[1]]
    first = {h: ws.cell(2, i + 1).value for i, h in enumerate(header)}
    assert (first["site_name"], first["original_site_name"], first["source_language"]) == ("THE FUNGUS", "EL HONGO", "es")
    assert "translation_note" in header


def test_the_readme_tab_lists_translation_note_as_editable(tmp_path):
    from waste_ai_search.workbook_io import REVIEW_INSTRUCTIONS

    queue = next(r for r in REVIEW_INSTRUCTIONS if r["tab_name"] == "Review_Queue")
    assert "translation_note" in queue["editable"]
    assert "original_site_name" in queue["sme_action"]
