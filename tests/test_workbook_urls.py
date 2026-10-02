"""Special characters in the review workbook: URLs and free text from the agent.

Prompted by the sibling refining-ai-search repo's broken-link work (RDP-63). Its finding was that
the pipeline copied URLs faithfully and the agent's own encodings were at fault. Our writer is
different code, so it was tested directly, and two real exposures turned up that RDP-63 did not
have:

- A control character in ANY cell made openpyxl raise IllegalCharacterError, aborting the whole
  workbook - every row of the run - after every agent call had been paid for.
- Link targets were written verbatim, so raw spaces, accents and non-ASCII hosts went into the
  workbook unencoded, and any scheme at all became clickable.

What these tests must NOT do is decode existing escapes: RDP-63 showed blanket decoding breaks
links that work.
"""
from __future__ import annotations

import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from waste_ai_search.workbook_io import (
    MAX_HYPERLINK_LENGTH,
    clean_cell_text,
    hyperlink_target,
    write_sheet,
)


# --- existing escapes are never touched (the RDP-63 lesson) ----------------------------------
@pytest.mark.parametrize(
    "url",
    [
        # Real URLs from this repo's cached responses, which work as given.
        "https://geocoder.ca/20%20Morrison%20Rd%2C%20Brantford%2C%20ON",
        "https://infofirma.sea.gob.cl/DocumentosSEA/MostrarDocumento?docId=86%2F53%2Fa13f905cf5288383f1d8e543258934a304",
        "https://cdm.unfccc.int/Projects/DB/DNV-CUK1166695034.41/view?cp=1",
        "https://example.org/report.pdf#page=4",
        "https://example.org/a?x=1&y=2;z=(3)",
    ],
)
def test_an_already_valid_url_is_left_exactly_as_it_is(url):
    assert hyperlink_target(url) == url


def test_percent_escapes_are_not_double_encoded():
    assert "%252C" not in hyperlink_target("https://example.org/a%2Cb")


# --- unsafe characters are encoded in the target -----------------------------------------
def test_a_raw_space_is_encoded():
    assert hyperlink_target("https://example.org/a file.pdf") == "https://example.org/a%20file.pdf"


def test_a_raw_accent_is_percent_encoded_as_utf8():
    """Local-language search (WP-480) makes accented paths more likely, e.g. Portuguese."""
    target = hyperlink_target("https://www.prefeitura.sp.gov.br/aterro-sanitário")
    assert target == "https://www.prefeitura.sp.gov.br/aterro-sanit%C3%A1rio"
    assert target.isascii()


def test_a_non_ascii_host_is_converted_to_idna():
    assert hyperlink_target("https://münchen.de/abfall") == "https://xn--mnchen-3ya.de/abfall"


def test_port_and_user_info_survive_host_encoding():
    assert hyperlink_target("https://user@münchen.de:8443/x") == "https://user@xn--mnchen-3ya.de:8443/x"


def test_an_ipv6_host_is_left_whole():
    assert hyperlink_target("https://[2001:db8::1]:8080/x") == "https://[2001:db8::1]:8080/x"


# --- what is never made clickable --------------------------------------------------------
@pytest.mark.parametrize(
    "url",
    ["javascript:alert(1)", "file:///C:/Windows/system32", "ftp://example.org/x", "mailto:a@b.org",
     "data:text/html,<script>", "example.org/no-scheme", "", None, "   "],
)
def test_only_http_and_https_become_links(url):
    """Agent output is untrusted; nothing but a web link should be clickable for an SME."""
    assert hyperlink_target(url) is None


def test_a_url_beyond_excels_limit_is_not_linked():
    assert hyperlink_target("https://example.org/" + "a" * MAX_HYPERLINK_LENGTH) is None


def test_a_url_at_the_limit_is_still_linked():
    url = "https://example.org/" + "a" * (MAX_HYPERLINK_LENGTH - len("https://example.org/"))
    assert len(url) == MAX_HYPERLINK_LENGTH
    assert hyperlink_target(url) == url


# --- control characters ------------------------------------------------------------------
@pytest.mark.parametrize("char", ["\x00", "\x08", "\x0b", "\x0c", "\x1f"])
def test_characters_openpyxl_rejects_become_spaces(char):
    assert clean_cell_text(f"page one{char}page two") == "page one page two"


@pytest.mark.parametrize("char", ["\t", "\n", "\r"])
def test_ordinary_whitespace_is_kept(char):
    assert clean_cell_text(f"a{char}b") == f"a{char}b"


def test_non_text_values_pass_through():
    assert clean_cell_text(12) == 12 and clean_cell_text(None) is None


def test_a_control_character_no_longer_aborts_the_workbook(tmp_path):
    """Before: IllegalCharacterError, and the run's whole review workbook was lost."""
    wb = Workbook()
    write_sheet(
        wb.active,
        ["evidence_summary", "winning_source_url"],
        [{"evidence_summary": "Quoted from a PDF:\x0cpage 2", "winning_source_url": "https://example.org/a\x0bb"}],
    )
    path = tmp_path / "w.xlsx"
    wb.save(path)

    ws = load_workbook(path).active
    assert ws["A2"].value == "Quoted from a PDF: page 2"


# --- end to end: what the workbook actually contains -------------------------------------
def test_the_workbook_shows_the_original_and_links_to_the_encoded_target(tmp_path):
    """The SME sees exactly what the source said; only the link target is made safe."""
    original = "https://www.prefeitura.sp.gov.br/aterro sanitário"
    wb = Workbook()
    write_sheet(wb.active, ["winning_source_url"], [{"winning_source_url": original}])
    path = tmp_path / "w.xlsx"
    wb.save(path)

    cell = load_workbook(path).active["A2"]
    assert cell.value == original
    assert cell.hyperlink.target == "https://www.prefeitura.sp.gov.br/aterro%20sanit%C3%A1rio"

    with zipfile.ZipFile(path) as z:
        rels = next(n for n in z.namelist() if "worksheets/_rels" in n)
        assert "aterro%20sanit%C3%A1rio" in z.read(rels).decode()


def test_an_unlinkable_url_stays_as_plain_visible_text(tmp_path):
    wb = Workbook()
    write_sheet(wb.active, ["url"], [{"url": "javascript:alert(1)"}])
    path = tmp_path / "w.xlsx"
    wb.save(path)

    cell = load_workbook(path).active["A2"]
    assert cell.value == "javascript:alert(1)"
    assert cell.hyperlink is None
