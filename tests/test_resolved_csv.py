"""Nothing is accepted automatically.

Every found value that differs from ours goes to the review queue, whatever its tier. A value equal
to ours is a confirmation: its evidence is kept, but it is neither queued nor loaded. resolved.csv
is written with its header and no rows; ingestion reads the SMEs' decisions from the workbook.

This reverses Q30's auto-validation of Tier 1-2 fills and of confirmations, on purpose: in batch 1,
2,932 values were accepted that way, including the 131 a map mirror won.
"""
from __future__ import annotations

import csv
import json

from waste_ai_search.arbitrate import run_arbitration
from waste_ai_search.cli import read_arbitrated_rows
from waste_ai_search.input_loader import write_csv_records
from waste_ai_search.run_context import PipelineConfig
from waste_ai_search.schema import RESOLVED_HEADERS, REVIEW_QUEUE_HEADERS
from waste_ai_search.seed_refresh import refresh_sites

REGULATOR = {
    "source_title": "Landfill permit",
    "publisher": "Environment Agency",
    "source_type": "Regulator",
    "publication_date": "2024-03-01",
    "url": "https://regulator.example.gov/permit/7",
}
NEWS = {
    "source_title": "Landfill expands",
    "publisher": "Local paper",
    "source_type": "News report",
    "publication_date": "2023-05-01",
    "url": "https://news.example.com/landfill",
}


def row(attribute, resolution, status, value="x"):
    return {
        "site_id": "1", "attribute_name": attribute, "resolution": resolution,
        "validation_status": status, "resolved_value": value, "winning_source_tier": "Tier 3",
    }


def found(name, value, source):
    return {
        "attribute_name": name, "value": value, "unit": "", "value_basis": "Direct",
        "confidence_score": "High", "value_date": "2024", "evidence_summary": "The permit says so.",
        "sources": [source],
    }


def arbitrate(tmp_path, *attributes, leftovers=()):
    """Arbitrate one OSM site, whose facility_type we already hold, against `attributes`.

    `leftovers` names files an earlier arbitration left in the run folder."""
    from waste_ai_search.seed_source import seed_headers

    run_dir = tmp_path / "run"
    site = {
        "facility_id": "7", "site_id": "7", "internal_facility_id": "7", "year": "2022",
        "site_name": "Seaton Meadows", "facility_name": "Seaton Meadows", "country_iso3": "GBR",
        "iso3c_plus": "GBR", "latitude": "54.6", "longitude": "-1.2", "facility_type": "Sanitary Landfill",
        "contributing_data_sources": "osm_2022", "reference_year": "2022",
    }
    write_csv_records(run_dir / "seed.csv", [site], seed_headers())
    raw = run_dir / "raw_foundry_responses"
    raw.mkdir(parents=True)
    (raw / "site_7.json").write_text(json.dumps({"site_id": "7", "attributes": list(attributes)}))
    for name in leftovers:
        (run_dir / name).write_text("left by an earlier arbitration")
    run_arbitration(PipelineConfig(input_csv=None, run_dir=run_dir, run_id="run", dataset_version="v"))
    return run_dir


def read(path):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def test_a_tier_1_fill_is_queued_and_resolved_csv_is_header_only(tmp_path):
    run_dir = arbitrate(
        tmp_path,
        found("facility_status", "Active", REGULATOR),
        found("opening_year", "1998", NEWS),
    )

    with (run_dir / "resolved.csv").open(encoding="utf-8-sig") as handle:
        assert handle.read().splitlines() == [",".join(RESOLVED_HEADERS)]

    queue = {r["attribute_name"]: r for r in read(run_dir / "review_queue.csv")}
    # A Tier 1 regulator filling an empty field once auto-validated; now an SME sees it.
    assert queue["facility_status"]["resolution"] == "Filled empty baseline"
    assert queue["facility_status"]["winning_source_tier"] == "Tier 1"
    assert queue["opening_year"]["winning_source_tier"] == "Tier 3"
    for queued in queue.values():
        assert queued["validation_status"] == "Needs review"
        assert (queued["reviewer"], queued["reviewed_date"]) == ("", "")


def test_no_standardized_table_or_load_statement_is_written(tmp_path):
    """Both were built before review; ingestion builds the AI search's table from reviewed values."""
    from openpyxl import load_workbook

    run_dir = arbitrate(tmp_path, found("facility_status", "Active", REGULATOR))

    assert not list(run_dir.glob("std_facility_tbl_*"))
    assert not list(run_dir.glob("*.load.sql"))
    settings = [row[0] for row in load_workbook(run_dir / "run_review.xlsx")["Run_Config"].iter_rows(values_only=True)]
    assert "standardized_records" not in settings


def test_arbitrating_again_removes_the_standardized_table_an_earlier_arbitration_left(tmp_path, capsys):
    """Batch 1 was arbitrated when both were still written, and its blob copy brings them back
    down. Left beside a header-only resolved.csv, they would look like values ready to load."""
    leftovers = ["std_facility_tbl_ai_search_run.csv", "std_facility_tbl_ai_search_run.load.sql"]
    run_dir = arbitrate(tmp_path, found("facility_status", "Active", REGULATOR), leftovers=leftovers)

    assert not list(run_dir.glob("std_facility_tbl_*"))
    assert not list(run_dir.glob("*.load.sql"))
    assert "Removed std_facility_tbl_ai_search_run.csv, std_facility_tbl_ai_search_run.load.sql" in capsys.readouterr().out
    assert [r["attribute_name"] for r in read(run_dir / "review_queue.csv")] == ["facility_status"]


def test_the_workbook_never_tells_smes_anything_was_accepted_for_them():
    from waste_ai_search.schema import DEFINITION_VALUES
    from waste_ai_search.workbook_io import REVIEW_INSTRUCTIONS

    text = " ".join(str(value) for tab in REVIEW_INSTRUCTIONS for value in tab.values())
    assert "auto-valid" not in text.lower()
    assert "AI Agent" not in text
    assert "std_facility" not in text
    assert "Auto-validated" not in DEFINITION_VALUES["validation_status"]


def test_a_confirmation_is_kept_as_evidence_but_not_queued(tmp_path):
    run_dir = arbitrate(tmp_path, found("facility_type", "sanitary landfill", REGULATOR))

    assert "facility_type" not in {r["attribute_name"] for r in read(run_dir / "review_queue.csv")}
    assert read(run_dir / "resolved.csv") == []
    evidence = [r for r in read(run_dir / "evidence.csv") if r["attribute_name"] == "facility_type"]
    assert [(r["mapped_value"], r["source_tier"]) for r in evidence] == [("Sanitary Landfill", "Tier 1")]


def test_attributes_not_found_stay_out_of_the_queue(tmp_path):
    run_dir = arbitrate(tmp_path, found("facility_status", "Active", REGULATOR))

    queued = {r["attribute_name"] for r in read(run_dir / "review_queue.csv")}
    assert queued == {"facility_status"}


def test_gas_collection_awaiting_review_still_unlocks_the_follow_up(tmp_path):
    # Pass 2 depends on this: a newly found gas collection system waits for review - nothing is
    # in resolved.csv - but must still reach the refreshed seed.
    write_csv_records(tmp_path / "resolved.csv", [], RESOLVED_HEADERS)
    queued = row("has_landfill_gas_collection", "Filled empty baseline", "Needs review", "TRUE")
    write_csv_records(tmp_path / "review_queue.csv", [queued], REVIEW_QUEUE_HEADERS)

    sites = [{"site_id": "1", "has_landfill_gas_collection": ""}]
    refreshed, stats = refresh_sites(sites, read_arbitrated_rows(tmp_path), "run")

    assert refreshed[0]["has_landfill_gas_collection"] == "TRUE"
    assert stats["gas_collection_discovered"] == 1


def test_a_tier_1_gas_collection_found_by_a_real_arbitration_unlocks_the_follow_up(tmp_path):
    """End to end: the finding is in the queue, not resolved.csv, and the refresh still reads it."""
    run_dir = arbitrate(tmp_path, found("has_landfill_gas_collection", "Yes", REGULATOR))

    sites = [{"site_id": "7", "has_landfill_gas_collection": ""}]
    refreshed, stats = refresh_sites(sites, read_arbitrated_rows(run_dir), "run")

    assert refreshed[0]["has_landfill_gas_collection"] == "TRUE"
    assert stats["gas_collection_discovered"] == 1
