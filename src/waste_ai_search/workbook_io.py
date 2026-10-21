from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any
from urllib.parse import quote, urlsplit, urlunsplit

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from .schema import (
    CORRECTED_VALUE_CHOICES,
    DEFINITION_VALUES,
    FIELD_TO_DEFINITION,
    PARSE_WARNING_HEADERS,
    REVIEW_QUEUE_HEADERS,
    SUPPLEMENTARY_LEADS_HEADERS,
)


# Excel will not follow a hyperlink longer than this; the text stays, the link is not made.
MAX_HYPERLINK_LENGTH = 2079

# Only these become clickable. Agent output is untrusted, and a `javascript:` or `file:` link must
# never turn into something an SME can click in the review workbook.
HYPERLINK_SCHEMES = {"http", "https"}

# Characters left as-is when encoding a link target. `%` is among them on purpose: existing escapes
# are never decoded or re-encoded. The sibling refining-ai-search repo (RDP-63) found that a
# blanket decode breaks links that work - most percent-encoded URLs there served fine as given.
_SAFE_PATH = "/:@!$&'()*+,;=%-._~"
_SAFE_QUERY = _SAFE_PATH + "?"


def clean_cell_text(value: Any) -> Any:
    """Make a value safe to write to a worksheet.

    openpyxl refuses control characters (vertical tab, form feed and the like) with
    IllegalCharacterError, and the error aborts the WHOLE workbook - every row of the run, after
    every agent call has been paid for. Quotes copied out of PDFs are where they come from. They are
    replaced with a space rather than dropped, so words either side do not run together.
    """
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub(" ", value)
    return value


def hyperlink_target(url: Any) -> str | None:
    """The link target to attach to a URL cell, or None if it should not be clickable.

    The cell keeps showing exactly what the source said; only the target is made safe. Raw spaces
    and non-ASCII characters are percent-encoded as UTF-8, and a non-ASCII host is converted to its
    ASCII (IDNA) form, which is what a browser does when the link is followed.
    """
    text = clean_cell_text(str(url).strip()) if url is not None else ""
    if not text or len(text) > MAX_HYPERLINK_LENGTH:
        return None
    try:
        parts = urlsplit(text)
    except ValueError:
        return None
    if parts.scheme.lower() not in HYPERLINK_SCHEMES or not parts.netloc:
        return None

    # Only the host is IDNA-encoded; user info and port are kept, and a bracketed IPv6 literal is
    # left whole rather than split on its colons.
    userinfo, at, hostport = parts.netloc.rpartition("@")
    if hostport.startswith("["):
        host, colon, port = hostport, "", ""
    else:
        host, colon, port = hostport.partition(":")
    if not host.isascii():
        try:
            host = host.encode("idna").decode("ascii")
        except UnicodeError:
            return None
    netloc = f"{userinfo}{at}{host}{colon}{port}"

    target = urlunsplit((
        parts.scheme,
        netloc,
        quote(parts.path, safe=_SAFE_PATH),
        quote(parts.query, safe=_SAFE_QUERY),
        quote(parts.fragment, safe=_SAFE_QUERY),
    ))
    return target if len(target) <= MAX_HYPERLINK_LENGTH else None


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
            header = ws.cell(1, cell.column).value
            if header and "url" in str(header).lower() and cell.value:
                target = hyperlink_target(cell.value)
                if target:
                    cell.hyperlink = target
                    cell.style = "Hyperlink"
    for idx, col in enumerate(ws.columns, start=1):
        header = ws.cell(1, idx).value or ""
        max_len = len(str(header))
        for cell in list(col)[1:80]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 80))
        width = min(max(max_len + 2, 10), 54)
        if any(term in str(header).lower() for term in ["summary", "notes", "reason", "url", "address"]):
            width = 54
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.auto_filter.ref = ws.dimensions


def write_sheet(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for record in rows:
        ws.append([clean_cell_text(record.get(header, "")) for header in headers])
        # openpyxl stores any string starting with "=" as a formula, and Excel evaluates it when the
        # workbook opens. Every value here is agent or source text, so none of it may be a formula:
        # "=HYPERLINK(...)" in a quote must be shown, not run.
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":
                cell.data_type = "s"
    style_sheet(ws)



REVIEW_INSTRUCTIONS = [
    {
        "tab_name": "Review_Queue",
        "purpose": (
            "YOUR WORKLIST. Every value the search found for an empty field, whatever the "
            "source's tier, and every value that disagrees with the database's from a Tier 1-3 "
            "source at least as credible as the database's own (for status, a Tier 1-3 source at "
            "least 3 years newer). Nothing is accepted without review. A weaker disagreement, a "
            "source whose link is dead and a unit that could not be converted are on the Leads tab "
            "instead. A source that only confirms the database's value is not shown, and nothing "
            "from it is loaded. Each value is filed under the site of ours its source is about, "
            "which may not be the site searched: every site column describes that site, "
            "searched_site_id is the search that found it, and match says why it is filed there. "
            "There is one row per site and attribute, and for a yearly attribute (status, waste "
            "received, waste in place, depth, cover types, gas figures) one per year, with every "
            "source filed under the site pooled into it. Ingestion reads your decisions from this tab."
        ),
        "sme_action": (
            "Work top to bottom. To see where the database puts the site, paste "
            "baseline_coordinates into Google Maps. Compare site_name with original_site_name - the source's own "
            "spelling - and note any translation problem in translation_note. Open "
            "winning_source_url and check the source says this, about THIS facility. A found_location row is one point, "
            "lat, lon, read from where its source says the site is (coordinates, a grid reference, or an address "
            "placed on its building; location_basis says which): paste it into Google Maps beside baseline_coordinates, and accept or reject the pair. "
            "Read match first: it says where the source places the facility and why the value is filed under this site. "
            "The match cell is coloured orange where the value is filed under a site other than the one searched, blue "
            "where its source gives no location (it stays with the site searched, so check the source is about this "
            "facility), and yellow where several of our sites tied. facility_name_in_source and stated_location are "
            "the facility as the source writes it. If the value is right but belongs to another of our sites, put that "
            "site's id in corrected_site_id (a dropdown where there were several candidates; any site id may be typed) "
            "and validate it, but first filter site_id to that site: if it already has a row for the same attribute "
            "(and, for a yearly attribute, the same year), reject one of the two, since ingestion takes one value per "
            "site, attribute and year. If it is about another of our sites and you can't tell which, reject it with 'Wrong "
            "facility - another site in our database'; if it is about a landfill we don't hold, with 'New facility - "
            "not in our database', which collect-new-facilities lists for ingestion. Set validation_status to Validated or Rejected; for a rejection, pick rejection_reason, and if you know the right value put it in corrected_value (a dropdown of allowed values for categorical fields; for a location, lat, lon) - never edit resolved_value. Each row of a yearly attribute is that year's value: the year of value_date, or this run's year where value_date is blank. So an older year's value is not Outdated: validate it if it was right for its year, and put a current figure on the newest year's row instead. If your corrected value is for a different date than value_date, put that date in corrected_value_date (YYYY, YYYY-MM or YYYY-MM-DD), but never a year that has its own row for the site and attribute. Put your name in "
            "reviewer and the date in reviewed_date, and say why in researcher_notes. An empty "
            "queue means nothing needs you."
        ),
        "editable": (
            "validation_status, rejection_reason, corrected_value, corrected_value_date, corrected_site_id, "
            "reviewer, reviewed_date, "
            "researcher_notes, "
            "translation_note - "
            "nothing else"
        ),
    },
    {
        "tab_name": "(file) evidence.csv",
        "purpose": (
            "One row per source-backed claim, with the tier rule applied to each source. Sources "
            "that only confirm the database's value are here too. Every row also gives the facility "
            "its source describes, by name and stated location, and the point read from that location; "
            "a source that states a location has a found_location row of its own."
        ),
        "sme_action": "Read only. Open it to see why a value won or lost.",
        "editable": "nothing",
    },
    {
        "tab_name": "(file) new_facility_candidates.csv",
        "purpose": (
            "Landfills none of our sites is: sources that place a facility more than 2 km from every site "
            "of ours, grouped where they lie within 2 km of each other, plus the values SMEs rejected as "
            "'New facility - not in our database' (added by collect-new-facilities). For ingestion."
        ),
        "sme_action": "Read only. Nothing from it is in Review_Queue.",
        "editable": "nothing",
    },
    {
        "tab_name": "(file) possible_duplicates.csv",
        "purpose": (
            "Pairs of our sites that may be one facility recorded twice: a source placed within 2 km of "
            "both that its name couldn't settle, or one page cited by both sites' searches. For ingestion."
        ),
        "sme_action": "Read only.",
        "editable": "nothing",
    },
    {
        "tab_name": "(file) datasets_found.csv",
        "purpose": (
            "Pages that look like datasets, such as a regulator's list of its landfills: naming or placing "
            "2 or more facilities, or from a publisher of datasets and found by 3 or more searches. A search "
            "reads only the entry it used, so ingestion can load the rest of the page once, as a source of its own."
        ),
        "sme_action": "Read only.",
        "editable": "nothing",
    },
    {
        "tab_name": "(file) outside_run_matches.csv",
        "purpose": (
            "Sources this run filed under a site outside its selection: another batch's country, or an "
            "excluded country it did not search. Their values are not in Review_Queue."
        ),
        "sme_action": "Read only.",
        "editable": "nothing",
    },
    {
        "tab_name": "Leads",
        "purpose": (
            "Findings that lost: a disagreement with the database's value from a source below the "
            "credibility floor (Tier 4-5) or weaker than the database's, a source whose link is "
            "dead, a unit that could not be converted, a 'Not a Waste Facility' verdict about a site it "
            "was not offered for; plus the agent's own unverified leads."
        ),
        "sme_action": "Read only. Promote something by hand if you judge the source good enough.",
        "editable": "nothing",
    },
    {
        "tab_name": "Parse_Warnings",
        "purpose": (
            "Values the pipeline dropped or could not map, including unit and enum problems, and "
            "citations of the two map sites that are never used as sources."
        ),
        "sme_action": "Scan for systematic problems worth fixing upstream.",
        "editable": "nothing",
    },
    {
        "tab_name": "Definitions",
        "purpose": "Allowed values, mirrored from the live database enums.",
        "sme_action": "Read only.",
        "editable": "nothing",
    },
]

RESOLUTION_FILLS = {
    "Conflict - needs review": "FFF2CC",
    "Overrides baseline": "FCE4D6",
    "Filled empty baseline": "E2EFDA",
    "Insufficient credibility": "EDEDED",
}


def add_definitions(ws) -> None:
    """One column per vocabulary, used as the source for the dropdowns."""
    names = [
        "review_decision",
        "rejection_reason",
        "resolution",
        "validation_status",
        "source_tier",
        "facility_status",
        "facility_type",
        "cover_type",
        "confidence_score",
        "bulk_waste_type",
        "waste_depth",
        "true_false",
        "gccs_energy_project_type",
        "gccs_current_project_status",
    ]
    ws.append(names)
    longest = max(len(DEFINITION_VALUES[name]) for name in names)
    for index in range(longest):
        ws.append([
            DEFINITION_VALUES[name][index] if index < len(DEFINITION_VALUES[name]) else ""
            for name in names
        ])
    style_sheet(ws)


def add_validations(wb, ws) -> None:
    """Attach dropdowns to the editable review columns only.

    A dropdown on a pipeline-set column (resolution, tiers) reads as an invitation to change it.
    """
    def_sheet = wb["Definitions"]
    columns = {def_sheet.cell(1, c).value: c for c in range(1, def_sheet.max_column + 1)}
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    last_row = max(ws.max_row, 2)

    for header, definition_name in FIELD_TO_DEFINITION.items():
        if header not in headers or definition_name not in columns:
            continue
        col = columns[definition_name]
        letter = get_column_letter(col)
        count = len(DEFINITION_VALUES[definition_name])
        formula = f"=Definitions!${letter}$2:${letter}${count + 1}"
        validation = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
        ws.add_data_validation(validation)
        target = get_column_letter(headers[header])
        validation.add(f"{target}2:{target}{last_row}")

    # corrected_value offers the allowed values of the attribute on its row, read from the same
    # Definitions columns as everything else; attributes without a vocabulary stay free text.
    if {"corrected_value", "attribute_name"} <= set(headers):
        rows_by_attribute: dict[str, list[int]] = {}
        for r in range(2, ws.max_row + 1):
            attribute = ws.cell(r, headers["attribute_name"]).value
            if attribute in CORRECTED_VALUE_CHOICES:
                rows_by_attribute.setdefault(attribute, []).append(r)
        target = get_column_letter(headers["corrected_value"])
        for attribute, rows in rows_by_attribute.items():
            definition_name, strict = CORRECTED_VALUE_CHOICES[attribute]
            if definition_name not in columns:
                continue
            letter = get_column_letter(columns[definition_name])
            count = len(DEFINITION_VALUES[definition_name])
            validation = DataValidation(
                type="list", formula1=f"=Definitions!${letter}$2:${letter}${count + 1}",
                allow_blank=True, showDropDown=False, showErrorMessage=strict,
            )
            ws.add_data_validation(validation)
            for start, end in _runs(rows):
                validation.add(f"{target}{start}" if start == end else f"{target}{start}:{target}{end}")


def _runs(rows: list[int]) -> list[tuple[int, int]]:
    """Consecutive row numbers as (first, last) runs, to keep a validation's range list short."""
    runs: list[tuple[int, int]] = []
    for r in rows:
        if runs and r == runs[-1][1] + 1:
            runs[-1] = (runs[-1][0], r)
        else:
            runs.append((r, r))
    return runs


def add_site_choices(ws, rows: list[dict[str, Any]]) -> None:
    """Offer each row's candidate sites in corrected_site_id, where there is more than one.

    The candidates are the site the value is filed under, the site searched, and the other sites
    its source fitted: our sites within 2 km of where it says the facility is, or namesakes of
    its name. For a source with no location they also include every other site whose search cited
    the same page. A typed site id is still accepted, so the list suggests rather than restricts.
    """
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    if "corrected_site_id" not in headers:
        return
    target = get_column_letter(headers["corrected_site_id"])
    rows_by_choices: dict[tuple[str, ...], list[int]] = {}
    for offset, record in enumerate(rows):
        choices = tuple(str(choice) for choice in record.get("_site_choices") or ())
        if len(choices) > 1:
            rows_by_choices.setdefault(choices, []).append(offset + 2)
    for choices, row_numbers in rows_by_choices.items():
        # A list written into the rule itself may hold at most 255 characters.
        listed = ",".join(choices)
        while len(listed) > 253 and "," in listed:
            listed = listed.rsplit(",", 1)[0]
        validation = DataValidation(
            type="list", formula1=f'"{listed}"', allow_blank=True, showDropDown=False, showErrorMessage=False,
        )
        ws.add_data_validation(validation)
        for start, end in _runs(row_numbers):
            validation.add(f"{target}{start}" if start == end else f"{target}{start}:{target}{end}")


# The match column is coloured where a value needs a second look at which site it belongs to: it
# is filed under a site other than the one searched, its source gives no location, or several of
# our sites tied for it.
MATCH_FILLS = {
    "another site": "F4B183",
    "no location": "BDD7EE",
    "tie": "FFD966",
}


def add_conditionals(ws) -> None:
    """Shade each row by its resolution so the queue reads at a glance, and the match column by
    why the value is filed where it is."""
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    last = max(ws.max_row, 2)
    if {"match", "site_id", "searched_site_id"} <= set(headers):
        # Added first, so on the match column they win over the row's resolution colour.
        match = get_column_letter(headers["match"])
        site = get_column_letter(headers["site_id"])
        searched = get_column_letter(headers["searched_site_id"])
        cells = f"{match}2:{match}{last}"
        formulas = {
            "another site": f'AND(${searched}2<>"",${site}2<>${searched}2)',
            "no location": f'LEFT(${match}2,11)="No location"',
            "tie": f'LEFT(${match}2,3)="Tie"',
        }
        for name, formula in formulas.items():
            ws.conditional_formatting.add(
                cells, FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=MATCH_FILLS[name]))
            )
    if "resolution" not in headers:
        return
    letter = get_column_letter(headers["resolution"])
    span = f"A2:{get_column_letter(ws.max_column)}{last}"
    for value, color in RESOLUTION_FILLS.items():
        ws.conditional_formatting.add(
            span,
            FormulaRule(formula=['$%s2="%s"' % (letter, value)], fill=PatternFill("solid", fgColor=color)),
        )


def write_review_workbook(
    path: Path,
    *,
    run_config: list[dict[str, Any]],
    review_queue: list[dict[str, Any]],
    leads: list[dict[str, Any]],
    parse_warnings: list[dict[str, Any]],
) -> Path:
    """The SME-facing artifact: only rows that need a human decision (Q31).

    Evidence and sources stay as CSV sidecars for audit. They are far too large to review by hand at
    full scale, and nothing in them requires a decision.
    """
    wb = Workbook()
    wb.remove(wb.active)

    ws_readme = wb.create_sheet("ReadMe")
    write_sheet(ws_readme, ["tab_name", "purpose", "sme_action", "editable"], REVIEW_INSTRUCTIONS)

    ws_config = wb.create_sheet("Run_Config")
    write_sheet(ws_config, ["setting", "value"], run_config)

    ws_queue = wb.create_sheet("Review_Queue")
    write_sheet(ws_queue, REVIEW_QUEUE_HEADERS, review_queue)


    ws_leads = wb.create_sheet("Leads")
    write_sheet(ws_leads, SUPPLEMENTARY_LEADS_HEADERS, leads)

    ws_warnings = wb.create_sheet("Parse_Warnings")
    write_sheet(ws_warnings, PARSE_WARNING_HEADERS, parse_warnings)

    ws_defs = wb.create_sheet("Definitions")
    add_definitions(ws_defs)

    add_validations(wb, ws_queue)
    add_site_choices(ws_queue, review_queue)
    add_conditionals(ws_queue)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
