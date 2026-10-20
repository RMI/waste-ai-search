"""Re-key a finished run's results from internal_facility_id to wastemap_id.

Runs from before RMI/waste_data_ingestion_pipeline#196 - batch 1 is the only one - used
`internal_facility_id` as every output's `site_id`. Each consolidation rebuild renumbered it, so
those IDs no longer name a facility. A facility's permanent ID is now its `wastemap_id`, and the
ingestion pipeline's raw load (RMI/waste_data_ingestion_pipeline#261) reads `site_id` as one.

The bridge is the source records. The run's own snapshot (consolidation_snapshot/, WP-548) holds
the crosswalk it searched against, which maps each old ID to its (data_source, facility_id)
records; today's `entity_linkage.crosswalk` maps the same records to wastemap_ids. Joining the two
gives each old ID the wastemap_ids its records now belong to:

- exactly one, which no other old ID in the run also gets: re-keyed;
- none: unmapped - none of its records is linked to a facility any more;
- more than one: split - its records now belong to different facilities;
- one that another of the run's old IDs also gets: merged - two searched sites are now one
  facility, and their results would collide in the ingestion transform (two values for one
  attribute, or one search's latitude paired with the other's longitude).

Any of the last three stops the relabel, and only the report is written: choosing a wastemap_id for
them is a judgement, not a lookup. An ID whose records all point at one wastemap_id is re-keyed even
when some of them are no longer linked, and the report lists it.

Only resolved.csv and the review workbook are rewritten - the two files the raw load reads - and
only into a new folder, never over the originals. Nothing changes but the IDs, which is checked by
reading the copies back. The workbook's Run_Config tab gains two rows saying what was done.
"""
from __future__ import annotations

import csv
import gzip
import os
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from itertools import zip_longest
from pathlib import Path
from typing import Any, Iterable, Mapping

from .schema import RESOLVED_HEADERS, REVIEW_QUEUE_HEADERS, normalize_scalar


# The crosswalk batch 1's consolidation consumed, as WP-548's snapshot saved it.
SNAPSHOT_CROSSWALK = Path("consolidation_snapshot") / "entity_linkage.crosswalk_legacy_replica_v1.csv.gz"
RESOLVED_NAME = "resolved.csv"
REPORT_NAME = "relabel_report.csv"
DEFAULT_OUT_DIR_NAME = "wastemap_id"

SITE_ID = "site_id"
# resolved.csv names the facility twice. Runs since #196 keep the internal_facility_id column but
# put the wastemap_id in it, so both are rewritten.
RESOLVED_ID_COLUMNS = (SITE_ID, "internal_facility_id")
REVIEW_QUEUE_SHEET = "Review_Queue"
RUN_CONFIG_SHEET = "Run_Config"
RUN_CONFIG_HEADERS = ["setting", "value"]
# The raw load's columns are RESOLVED_HEADERS and REVIEW_QUEUE_HEADERS, except that it also takes
# corrected_value_date, which a workbook from before WP-552 lacks and which then loads empty.
OPTIONAL_REVIEW_QUEUE_COLUMNS = {"corrected_value_date"}

# The Run_Config rows a relabel adds. Finding the first already there means the workbook has been
# relabelled, and relabelling it again would read wastemap_ids as old IDs.
SITE_ID_KEY_SETTING = "site_id_key"
RELABELLED_SETTING = "site_id_relabelled"

MAPPED = "mapped"
PARTLY_LINKED = "partly_linked"
UNMAPPED = "unmapped"
SPLIT = "split"
MERGED = "merged"
STOPPING_OUTCOMES = (UNMAPPED, SPLIT, MERGED)

REPORT_HEADERS = [
    "old_internal_facility_id", "outcome", "wastemap_id", "candidate_wastemap_ids",
    "source_records", "linked_records", "appears_in", "detail",
]

# How the database's crosswalk is chosen: the consolidation run a seed reads, as
# consolidation_run.py picks it in #22 - pinned by CONSOLIDATION_RUN_ID, or else the newest run of
# these strategies - and the crosswalk run that consolidation consumed.
DEFAULT_CONSOLIDATION_STRATEGY = "legacy_replica_v1"
DEFAULT_ENTITY_LINKAGE_STRATEGY = "legacy_replica_v1"
RUNS_SQL = (
    "SELECT consolidation_run_id, consolidation_strategy, entity_linkage_strategy,\n"
    "       MAX(entity_linkage_run_id) AS entity_linkage_run_id,\n"
    "       MAX(consolidation_created_date) AS created\n"
    "FROM consolidation.consolidated_facility\n"
    "GROUP BY consolidation_run_id, consolidation_strategy, entity_linkage_strategy\n"
    "ORDER BY created DESC NULLS LAST, consolidation_run_id DESC"
)
CROSSWALK_SQL = (
    "SELECT wastemap_id, data_source, facility_id\n"
    "FROM entity_linkage.crosswalk\n"
    "WHERE entity_linkage_run_id = %s"
)

Record = tuple[str, str]  # (data_source, facility_id)


class RelabelError(ValueError):
    """The inputs can't be relabelled as they are; the message says why and what to do."""


@dataclass
class Crosswalk:
    """Today's crosswalk: the wastemap_id(s) each source record belongs to."""

    wastemap_ids: dict[Record, set[str]]
    description: str


@dataclass
class IdMapping:
    old_id: str
    outcome: str
    # The ID the old one is rewritten to; empty unless the outcome is mapped or partly_linked.
    wastemap_id: str = ""
    candidates: list[str] = field(default_factory=list)
    source_records: int = 0
    linked_records: int = 0
    detail: str = ""


@dataclass
class RelabelResult:
    mappings: dict[str, IdMapping]
    report: Path
    written: list[Path]

    @property
    def stopped(self) -> list[IdMapping]:
        return [mapping for mapping in self.mappings.values() if mapping.outcome in STOPPING_OUTCOMES]


def _id_text(value: Any) -> str:
    """An ID as text. A spreadsheet can hand back 2514 as 2514.0, and an export as "2514.0"."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = normalize_scalar(value)
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def _id_order(old_id: str) -> tuple[int, int, str]:
    return (0, int(old_id), "") if old_id.isdigit() else (1, 0, old_id)


def _records_text(records: Iterable[Record]) -> str:
    return ", ".join(f"{data_source}:{facility_id}" for data_source, facility_id in records)


def _open_csv(path: Path):
    if path.suffix == ".gz":
        return gzip.open(path, "rt", newline="", encoding="utf-8")
    return path.open(newline="", encoding="utf-8")


def _require_columns(fieldnames: Iterable[str] | None, required: Iterable[str], path: Path) -> None:
    missing = [column for column in required if column not in (fieldnames or [])]
    if missing:
        raise RelabelError(f"{path} has no {', '.join(missing)} column; it has {list(fieldnames or [])}.")


# --- the two crosswalks -----------------------------------------------------------------------------

def read_snapshot_crosswalk(path: Path) -> dict[str, set[Record]]:
    """The run's own crosswalk: each old internal_facility_id's source records."""
    records: dict[str, set[Record]] = defaultdict(set)
    with _open_csv(path) as handle:
        reader = csv.DictReader(handle)
        _require_columns(reader.fieldnames, ("internal_facility_id", "data_source", "facility_id"), path)
        for row in reader:
            records[_id_text(row["internal_facility_id"])].add(_record(row))
    return dict(records)


def _record(row: Mapping[str, Any]) -> Record:
    return normalize_scalar(row["data_source"]), normalize_scalar(row["facility_id"])


def _crosswalk(rows: Iterable[Mapping[str, Any]], description: str) -> Crosswalk:
    wastemap_ids: dict[Record, set[str]] = defaultdict(set)
    for row in rows:
        wastemap_id = _id_text(row["wastemap_id"])
        if not wastemap_id:
            raise RelabelError(f"{description} has a row with no wastemap_id: {dict(row)}")
        wastemap_ids[_record(row)].add(wastemap_id)
    if not wastemap_ids:
        raise RelabelError(f"{description} is empty.")
    return Crosswalk(dict(wastemap_ids), description)


def read_crosswalk_csv(path: Path) -> Crosswalk:
    """Today's crosswalk from an export of entity_linkage.crosswalk (optionally gzipped).

    It needs wastemap_id, data_source and facility_id. The table keeps every entity linkage run side
    by side, so an export holding more than one run is refused: it would blend them.
    """
    path = Path(path)
    with _open_csv(path) as handle:
        reader = csv.DictReader(handle)
        _require_columns(reader.fieldnames, ("wastemap_id", "data_source", "facility_id"), path)
        has_runs = "entity_linkage_run_id" in reader.fieldnames
        rows = list(reader)
    runs = sorted({row["entity_linkage_run_id"] for row in rows}) if has_runs else []
    if len(runs) > 1:
        raise RelabelError(
            f"{path} holds {len(runs)} entity linkage runs ({', '.join(runs[:5])}). Export only the run "
            "the consolidation consumed; the README gives the query."
        )
    description = f"{path.name}" + (f" (entity_linkage_run_id {runs[0]})" if runs else "")
    return _crosswalk(rows, description)


def choose_consolidation_run(
    runs: list[Mapping[str, Any]], pinned: str = "", environ: Mapping[str, str] | None = None
) -> dict[str, str]:
    """Pick the consolidation run a seed reads from RUNS_SQL's rows (newest first)."""
    environ = os.environ if environ is None else environ
    if not runs:
        raise RelabelError("consolidation.consolidated_facility is empty, or this is the wrong database.")

    def describe(row: Mapping[str, Any], chosen_by: str) -> dict[str, str]:
        chosen = {
            key: normalize_scalar(row.get(key))
            for key in ("consolidation_run_id", "entity_linkage_run_id", "created")
        }
        if not chosen["entity_linkage_run_id"]:
            raise RelabelError(
                f"Consolidation run {chosen['consolidation_run_id']} records no entity_linkage_run_id, "
                "so the crosswalk it consumed can't be found. Pin another with --consolidation-run-id."
            )
        return {**chosen, "chosen_by": chosen_by}

    pinned = (pinned or environ.get("CONSOLIDATION_RUN_ID") or "").strip()
    if pinned:
        match = [row for row in runs if normalize_scalar(row.get("consolidation_run_id")) == pinned]
        if not match:
            raise RelabelError(f"Consolidation run {pinned!r} is not in consolidation.consolidated_facility.")
        return describe(match[0], "pinned")

    strategy = (environ.get("CONSOLIDATION_STRATEGY") or DEFAULT_CONSOLIDATION_STRATEGY).strip()
    linkage = (environ.get("ENTITY_LINKAGE_STRATEGY") or DEFAULT_ENTITY_LINKAGE_STRATEGY).strip()
    eligible = [
        row
        for row in runs
        if normalize_scalar(row.get("consolidation_strategy")) == strategy
        and normalize_scalar(row.get("entity_linkage_strategy")) == linkage
    ]
    if not eligible:
        raise RelabelError(
            f"No consolidation run for strategy {strategy}/{linkage}. Pin one with --consolidation-run-id."
        )
    return describe(eligible[0], f"the newest {strategy}/{linkage} run")


def read_crosswalk_db(config: Any = None, consolidation_run_id: str = "") -> Crosswalk:
    """Today's crosswalk from the database: the crosswalk run the chosen consolidation run consumed."""
    from .db import fetch_all

    run = choose_consolidation_run(fetch_all(RUNS_SQL, config=config), consolidation_run_id)
    rows = fetch_all(CROSSWALK_SQL, (run["entity_linkage_run_id"],), config=config)
    return _crosswalk(
        rows,
        f"entity_linkage.crosswalk run {run['entity_linkage_run_id']}, consumed by consolidation run "
        f"{run['consolidation_run_id']} ({run['chosen_by']}, created {run['created'] or 'at an unknown time'})",
    )


# --- mapping ------------------------------------------------------------------------------------------

def map_old_ids(
    old_ids: Iterable[str], snapshot_crosswalk: Mapping[str, set[Record]], crosswalk: Crosswalk
) -> dict[str, IdMapping]:
    """Each old ID's wastemap_id, through the source records the two crosswalks share."""
    mappings: dict[str, IdMapping] = {}
    for old_id in sorted(set(old_ids), key=_id_order):
        records = sorted(snapshot_crosswalk.get(old_id, ()))
        if not records:
            mappings[old_id] = IdMapping(old_id, UNMAPPED, detail="not in the run's own crosswalk")
            continue

        by_target: dict[str, list[Record]] = defaultdict(list)
        unlinked = []
        for record in records:
            targets = crosswalk.wastemap_ids.get(record)
            if not targets:
                unlinked.append(record)
            for target in sorted(targets or (), key=_id_order):
                by_target[target].append(record)
        mapping = IdMapping(
            old_id,
            UNMAPPED,
            candidates=sorted(by_target, key=_id_order),
            source_records=len(records),
            linked_records=len(records) - len(unlinked),
        )
        if not by_target:
            mapping.detail = f"none of its source records is in the current crosswalk: {_records_text(records)}"
        elif len(by_target) > 1:
            mapping.outcome = SPLIT
            mapping.detail = "its records now belong to different facilities: " + "; ".join(
                f"{target} <- {_records_text(by_target[target])}" for target in mapping.candidates
            )
        else:
            mapping.wastemap_id = mapping.candidates[0]
            mapping.outcome = PARTLY_LINKED if unlinked else MAPPED
            if unlinked:
                mapping.detail = (
                    f"re-keyed by {mapping.linked_records} of {len(records)} records; "
                    f"no longer linked: {_records_text(unlinked)}"
                )
        mappings[old_id] = mapping

    sharing: dict[str, list[str]] = defaultdict(list)
    for mapping in mappings.values():
        if mapping.wastemap_id:
            sharing[mapping.wastemap_id].append(mapping.old_id)
    for wastemap_id, old_ids_sharing in sharing.items():
        if len(old_ids_sharing) < 2:
            continue
        for old_id in old_ids_sharing:
            mapping = mappings[old_id]
            others = ", ".join(other for other in old_ids_sharing if other != old_id)
            note = f"wastemap_id {wastemap_id} is also the facility of old ID(s) {others}"
            mapping.detail = f"{note}; {mapping.detail}" if mapping.detail else note
            mapping.outcome = MERGED
            mapping.wastemap_id = ""
    return mappings


# --- the run's files ----------------------------------------------------------------------------------

def _header(sheet) -> list[str]:
    return [normalize_scalar(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]


def _site_id_column(sheet) -> int | None:
    header = _header(sheet)
    return header.index(SITE_ID) + 1 if SITE_ID in header else None


def _unnamed_columns_with_values(sheet) -> list[int]:
    """Columns with no header but some value, which pandas reads as 'Unnamed: N'."""
    header = _header(sheet)
    unnamed = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        for index, value in enumerate(row):
            if value not in (None, "") and (index >= len(header) or not header[index]):
                unnamed.add(index + 1)
    return sorted(unnamed)


def check_layout(resolved_header: list[str], workbook) -> None:
    """Stop unless the files have exactly the columns the raw load accepts, so a mismatch shows up
    here rather than at load time."""
    problems = []

    def compare(name: str, actual: list[str], expected: set[str], optional: set[str] = frozenset()) -> None:
        missing = sorted(expected - optional - set(actual))
        unexpected = sorted(set(actual) - expected)
        repeated = sorted(column for column, count in Counter(actual).items() if count > 1)
        if missing or unexpected or repeated:
            problems.append(f"{name}: missing {missing}, unexpected {unexpected}, repeated {repeated}")

    compare(RESOLVED_NAME, resolved_header, set(RESOLVED_HEADERS))
    expected_tabs = {
        REVIEW_QUEUE_SHEET: (set(REVIEW_QUEUE_HEADERS) | OPTIONAL_REVIEW_QUEUE_COLUMNS, OPTIONAL_REVIEW_QUEUE_COLUMNS),
        RUN_CONFIG_SHEET: (set(RUN_CONFIG_HEADERS), set()),
    }
    for name, (expected, optional) in expected_tabs.items():
        if name not in workbook.sheetnames:
            problems.append(f"the workbook has no {name} tab")
            continue
        compare(name, [column for column in _header(workbook[name]) if column], expected, optional)
        unnamed = _unnamed_columns_with_values(workbook[name])
        if unnamed:
            problems.append(f"{name}: values under columns with no header (columns {unnamed})")
    if problems:
        raise RelabelError(
            "The files don't have the columns the raw load (blob_to_db_raw_ai_search.py) expects, so "
            "they would not load even once re-keyed. " + " | ".join(problems)
        )


def run_config(workbook) -> dict[str, str]:
    rows = workbook[RUN_CONFIG_SHEET].iter_rows(min_row=2, values_only=True)
    return {normalize_scalar(row[0]): normalize_scalar(row[1] if len(row) > 1 else None) for row in rows if row and row[0]}


def read_resolved(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        return header, list(reader)


def site_ids_in(resolved_header: list[str], resolved_rows: list[list[str]], workbook) -> dict[str, Counter]:
    """Every old ID in the files, and how many rows of each file or tab it appears in."""
    found: dict[str, Counter] = defaultdict(Counter)
    site, internal = (resolved_header.index(column) for column in RESOLVED_ID_COLUMNS)
    differing = []
    for number, row in enumerate(resolved_rows, start=2):
        if _id_text(row[site]) != _id_text(row[internal]):
            differing.append(f"line {number}: {row[site]!r} vs {row[internal]!r}")
        if _id_text(row[site]):
            found[_id_text(row[site])][RESOLVED_NAME] += 1
    if differing:
        raise RelabelError(
            f"{RESOLVED_NAME} has rows whose site_id and internal_facility_id differ, so which one names "
            f"the facility is unclear: {'; '.join(differing[:5])}"
        )
    for sheet in workbook.worksheets:
        column = _site_id_column(sheet)
        if column is None:
            continue
        for (value,) in sheet.iter_rows(min_row=2, min_col=column, max_col=column, values_only=True):
            if _id_text(value):
                found[_id_text(value)][sheet.title] += 1
    return dict(found)


def _appears_in(counts: Counter) -> str:
    return "; ".join(f"{place}: {rows}" for place, rows in counts.items())


def write_report(path: Path, mappings: Mapping[str, IdMapping], where: Mapping[str, Counter]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(REPORT_HEADERS)
        # The IDs that stop the relabel come first.
        ordered = sorted(mappings.values(), key=lambda m: (m.outcome not in STOPPING_OUTCOMES, m.outcome, _id_order(m.old_id)))
        for mapping in ordered:
            writer.writerow([
                mapping.old_id, mapping.outcome, mapping.wastemap_id, "; ".join(mapping.candidates),
                mapping.source_records, mapping.linked_records, _appears_in(where.get(mapping.old_id, Counter())),
                mapping.detail,
            ])


def _temporary(path: Path) -> Path:
    """Hidden, and keeping its extension, which openpyxl needs to read it back."""
    return path.with_name(f".{path.stem}.tmp{path.suffix}")


def write_resolved(path: Path, header: list[str], rows: list[list[str]], id_map: Mapping[str, str]) -> None:
    """Written as the search writes it (csv's default dialect), so only the ID cells differ."""
    columns = [header.index(column) for column in RESOLVED_ID_COLUMNS]
    with _temporary(path).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for row in rows:
            row = list(row)
            for column in columns:
                if _id_text(row[column]):
                    row[column] = id_map[_id_text(row[column])]
            writer.writerow(row)


def write_workbook(workbook, path: Path, id_map: Mapping[str, str], provenance: list[tuple[str, str]]) -> None:
    """Rewrite every tab's site_id column in place, keeping each cell's type, and add provenance."""
    for sheet in workbook.worksheets:
        column = _site_id_column(sheet)
        if column is None:
            continue
        for (cell,) in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            old_id = _id_text(cell.value)
            if not old_id:
                continue
            new_id = id_map[old_id]
            numeric = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            cell.value = int(new_id) if numeric and new_id.isdigit() else new_id
    settings = workbook[RUN_CONFIG_SHEET]
    for row in provenance:
        settings.append(list(row))
    if settings.auto_filter.ref:
        settings.auto_filter.ref = f"A1:{settings.cell(1, settings.max_column).column_letter}{settings.max_row}"
    workbook.save(_temporary(path))


def _comparable(row: tuple, column: int | None, id_map: Mapping[str, str] | None = None) -> tuple:
    """A sheet row without trailing empty cells, its site_id as text - mapped, if id_map is given."""
    row = tuple(row)
    while row and row[-1] is None:
        row = row[:-1]
    if column is None or len(row) <= column or not _id_text(row[column]):
        return row
    site_id = _id_text(row[column])
    return row[:column] + (id_map[site_id] if id_map is not None else site_id,) + row[column + 1:]


def verify_copies(
    resolved: tuple[list[str], list[list[str]]],
    resolved_copy: Path,
    workbook_path: Path,
    workbook_copy: Path,
    id_map: Mapping[str, str],
    provenance: list[tuple[str, str]],
) -> None:
    """Read the copies back as the raw load would, and check nothing but the IDs changed."""
    import openpyxl

    problems = []
    header, rows = resolved
    copy_header, copy_rows = read_resolved(resolved_copy)
    columns = {header.index(column) for column in RESOLVED_ID_COLUMNS}
    if copy_header != header or len(copy_rows) != len(rows):
        problems.append(f"{RESOLVED_NAME}: header or row count changed")
    for number, (row, copy) in enumerate(zip(rows, copy_rows), start=2):
        expected = [id_map.get(_id_text(value), value) if index in columns else value for index, value in enumerate(row)]
        if copy != expected:
            problems.append(f"{RESOLVED_NAME} line {number} changed beyond its IDs")

    original = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    rewritten = openpyxl.load_workbook(workbook_copy, read_only=True, data_only=True)
    try:
        if original.sheetnames != rewritten.sheetnames:
            problems.append(f"tabs changed: {original.sheetnames} -> {rewritten.sheetnames}")
        for name in original.sheetnames:
            before = list(original[name].iter_rows(values_only=True))
            after = list(rewritten[name].iter_rows(values_only=True))
            if name == RUN_CONFIG_SHEET:
                if [_comparable(row, None) for row in after[len(before):]] != [tuple(row) for row in provenance]:
                    problems.append(f"{name}: the provenance rows are not what was written")
                after = after[: len(before)]
            header = _comparable(before[0], None) if before else ()
            column = header.index(SITE_ID) if SITE_ID in header else None
            for number, (row, copy) in enumerate(zip_longest(before, after, fillvalue=()), start=1):
                if number == 1:
                    changed = _comparable(row, None) != _comparable(copy, None)
                else:
                    changed = _comparable(row, column, id_map) != _comparable(copy, column)
                if changed:
                    problems.append(f"{name} row {number} changed beyond its site_id")
                    break
    finally:
        original.close()
        rewritten.close()
    if problems:
        raise RelabelError("The copies don't match the originals: " + "; ".join(problems[:10]))


def relabel_run(
    run_dir: Path,
    crosswalk: Crosswalk,
    out_dir: Path | None = None,
    workbook_path: Path | None = None,
    snapshot_crosswalk_path: Path | None = None,
    dry_run: bool = False,
    today: date | None = None,
) -> RelabelResult:
    """Re-key a run's resolved.csv and review workbook to wastemap_id, into out_dir.

    The workbook defaults to the run's own <run_id>_review.xlsx; pass the SMEs' copy downloaded from
    SharePoint to carry their decisions across. The report is always written; the copies only when
    every ID maps cleanly and dry_run is off.
    """
    import openpyxl

    run_dir = Path(run_dir).expanduser().resolve()
    out_dir = Path(out_dir).expanduser().resolve() if out_dir else run_dir / DEFAULT_OUT_DIR_NAME
    workbook_path = Path(workbook_path).expanduser().resolve() if workbook_path else run_dir / f"{run_dir.name}_review.xlsx"
    snapshot_crosswalk_path = Path(snapshot_crosswalk_path or run_dir / SNAPSHOT_CROSSWALK).expanduser().resolve()
    resolved_path = run_dir / RESOLVED_NAME
    for path in (resolved_path, workbook_path, snapshot_crosswalk_path):
        if not path.is_file():
            raise RelabelError(f"{path} does not exist.")
    if out_dir in (run_dir, workbook_path.parent):
        raise RelabelError(f"Write the copies to a folder of their own, not {out_dir}, which holds the originals.")

    # Formulas are read as their last calculated values - what the raw load reads - so the copy
    # holds values. data_only=False would keep the formulas but lose the values pandas reads.
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    resolved_header, resolved_rows = read_resolved(resolved_path)
    check_layout(resolved_header, workbook)
    settings = run_config(workbook)
    run_id = settings.get("run_id", "")
    if not run_id:
        raise RelabelError(f"{workbook_path.name}'s Run_Config tab records no run_id.")
    if SITE_ID_KEY_SETTING in settings:
        raise RelabelError(
            f"{workbook_path.name} already records {SITE_ID_KEY_SETTING} = {settings[SITE_ID_KEY_SETTING]!r}: "
            "it has been relabelled. Relabel the original instead."
        )

    outputs = [out_dir / RESOLVED_NAME, out_dir / f"{run_id}_review.xlsx"]
    existing = [path for path in outputs if path.exists()]
    if existing and not dry_run:
        raise RelabelError(
            f"{', '.join(str(path) for path in existing)} already exists. Choose an empty --out-dir; "
            "nothing is ever written over a run's files."
        )

    where = site_ids_in(resolved_header, resolved_rows, workbook)
    mappings = map_old_ids(where, read_snapshot_crosswalk(snapshot_crosswalk_path), crosswalk)
    report = out_dir / REPORT_NAME
    write_report(report, mappings, where)
    result = RelabelResult(mappings, report, [])
    if result.stopped or dry_run:
        return result

    id_map = {old_id: mapping.wastemap_id for old_id, mapping in mappings.items()}
    provenance = [
        (SITE_ID_KEY_SETTING, "wastemap_id"),
        (
            RELABELLED_SETTING,
            f"{(today or date.today()).isoformat()}: site_id was internal_facility_id in "
            f"{snapshot_crosswalk_path.name}; re-keyed through {crosswalk.description}. "
            f"Per-ID detail in {REPORT_NAME}.",
        ),
    ]
    try:
        write_resolved(outputs[0], resolved_header, resolved_rows, id_map)
        write_workbook(workbook, outputs[1], id_map, provenance)
        verify_copies(
            (resolved_header, resolved_rows), _temporary(outputs[0]), workbook_path, _temporary(outputs[1]), id_map, provenance
        )
        for path in outputs:
            os.replace(_temporary(path), path)
    finally:
        for path in outputs:
            _temporary(path).unlink(missing_ok=True)
    result.written = outputs
    return result
