"""What finished search runs found, by how searchable each site's name is.

Step 1 of docs/plans/2026-10-15-skip-unsearchable-names.md. Before any class of name is skipped,
this measures what searching it actually found, compared with named sites from the same kind of
source. Most unnamed sites come only from OpenStreetMap or Global Plastic Watch and start with an
empty record, so comparing them with all named sites would flatter the named ones; many of those
come from government registries holding far more data.

Read-only. A run's files come from blob storage straight into memory, or from a run folder on
disk, and SME decisions from the review workbook in the synced SharePoint folder, opened without
saving. Nothing is re-arbitrated, republished or written back.
"""
from __future__ import annotations

import csv
import io
import json
from collections import Counter, defaultdict
from contextlib import redirect_stdout
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Optional

from .schema import CONTRADICTION_CHECK_SOURCES, LAST_BATCH, NOT_A_WASTE_FACILITY, normalize_scalar
from .searchable_names import NAME_CLASSES, NAMED, name_class, name_reason
from .snapshot import MANIFEST_NAME, SNAPSHOT_DIR_NAME, describe_manifest_run


# A path inside a run folder -> its bytes, or None when the run has no such file.
Reader = Callable[[str], Optional[bytes]]

SEED = "seed.csv"
RUN_LOG = "foundry_run_log.csv"
RESOLVED = "resolved.csv"
REVIEW_QUEUE = "review_queue.csv"
MANIFEST = f"{SNAPSHOT_DIR_NAME}/{MANIFEST_NAME}"
REVIEW_SHEET = "Review_Queue"

MAP_ONLY = "OSM or Global Plastic Watch only"
OTHER_SOURCES = "has another source"
SOURCE_GROUPS = [MAP_ONLY, OTHER_SOURCES]

# What counts as a new value: the search filled an empty field or challenged an existing one.
# resolved.csv holds the values a run accepted without review. Only batch 1 and earlier runs have
# any: later runs accept nothing automatically, write it header-only and queue every new value. A
# confirmation adds nothing new: every row in the review queue is a candidate except one, which
# batch 1 queued only for a failed identity check and later runs never queue.
CONFIRMED = "Confirmed baseline"
NEW_VALUE_RESOLUTIONS = {"Filled empty baseline", "Overrides baseline"}
# An SME's rejection of a value as about another facility. Since build step 5 of the
# source-matching plan it says which kind ("Wrong facility - another site in our database"), so
# reasons are counted by how they start.
WRONG_FACILITY = "Wrong facility"
NAME_ATTRIBUTE = "found_facility_name"
# Run-log statuses for a site the agent actually answered; "Failed" and a failed search tool did not.
ANSWERED_STATUSES = ("Succeeded", "Skipped existing response")

# The bar in the plan: skip a class when its sites get a new value less than half as often as named
# sites from the same kind of source, or when SMEs rejected most of its reviewed values as the wrong
# facility. Fewer answered sites than this is too few to judge.
BAR = 0.5
MIN_SITES_TO_JUDGE = 30
# The rate behind the README's batch estimates: about $525 for 1,931 sites at about 19 queries each.
DEFAULT_COST_PER_1000_QUERIES = 14.0


def local_reader(run_dir: Path) -> Reader:
    def read(name: str) -> Optional[bytes]:
        path = Path(run_dir) / name
        return path.read_bytes() if path.is_file() else None

    return read


def blob_reader(client: Any, run_id: str) -> Reader:
    """Read a run's files from blob into memory; nothing is written to disk."""
    from azure.core.exceptions import ResourceNotFoundError

    from .storage import run_prefix

    def read(name: str) -> Optional[bytes]:
        try:
            return client.download_blob(f"{run_prefix(run_id)}/{name}").readall()
        except ResourceNotFoundError:
            return None

    return read


def list_blob_runs(client: Any) -> list[str]:
    """The run ids in blob, from the folder names alone, without listing every file in them."""
    prefix = "outputs/runs/"
    return sorted(
        item.name[len(prefix):].strip("/")
        for item in client.walk_blobs(name_starts_with=prefix, delimiter="/")
    )


def csv_rows(data: Optional[bytes]) -> list[dict[str, str]]:
    if data is None:
        return []
    return [dict(row) for row in csv.DictReader(io.StringIO(data.decode("utf-8-sig")))]


@dataclass
class RunFiles:
    run_id: str
    seed: list[dict[str, str]]
    run_log: list[dict[str, str]]
    resolved: list[dict[str, str]]
    review: list[dict[str, str]]
    manifest: Optional[dict[str, Any]]
    # Review_Queue rows from the reviewed workbook; None when there is no workbook to read.
    sme: Optional[list[dict[str, str]]] = None


def load_run(run_id: str, read: Reader) -> RunFiles:
    data = {name: read(name) for name in (SEED, RUN_LOG, RESOLVED, REVIEW_QUEUE)}
    missing = [name for name, content in data.items() if content is None]
    if missing:
        raise FileNotFoundError(
            f"Run {run_id} has no {', '.join(missing)}. A run must be searched and arbitrated to be "
            "measured."
        )
    manifest = read(MANIFEST)
    return RunFiles(
        run_id=run_id,
        seed=csv_rows(data[SEED]),
        run_log=csv_rows(data[RUN_LOG]),
        resolved=csv_rows(data[RESOLVED]),
        review=csv_rows(data[REVIEW_QUEUE]),
        manifest=json.loads(manifest) if manifest is not None else None,
    )


def read_review_workbook(path: Path) -> Optional[list[dict[str, str]]]:
    """The SMEs' Review_Queue rows, or None when the workbook isn't there. Opened read-only."""
    if not Path(path).is_file():
        return None
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if REVIEW_SHEET not in workbook.sheetnames:
            return None
        rows = workbook[REVIEW_SHEET].iter_rows(values_only=True)
        headers = [normalize_scalar(value) for value in next(rows, ())]
        return [
            dict(zip(headers, (normalize_scalar(value) for value in row)))
            for row in rows
            if any(value not in (None, "") for value in row)
        ]
    finally:
        workbook.close()


def source_group(site: dict[str, Any]) -> str:
    sources = {
        token.strip().lower()
        for token in normalize_scalar(site.get("contributing_data_sources")).split("+")
        if token.strip()
    }
    return MAP_ONLY if sources and sources <= CONTRADICTION_CHECK_SOURCES else OTHER_SOURCES


def _site_id(row: dict[str, Any]) -> str:
    return normalize_scalar(row.get("site_id"))


def _searched_id(row: dict[str, Any]) -> str:
    """The site whose search found a value. Since build step 5 a value may be filed under another
    site, and searched_site_id says which search found it; earlier runs file it where it was found."""
    return normalize_scalar(row.get("searched_site_id")) or _site_id(row)


def _as_int(value: Any) -> int:
    text = normalize_scalar(value)
    return int(text) if text.isdigit() else 0


SITE_HEADERS = [
    "run_id", "site_id", "country_iso3", "site_name", "original_site_name", "name_class",
    "source_group", "status", "answered", "web_searches", "new_values_accepted",
    "new_values_to_review", "found_attributes", "found_name", "filed_under_another_site",
    "not_a_waste_facility", "sme_validated", "sme_rejected", "sme_wrong_facility",
]


def site_rows(run: RunFiles) -> list[dict[str, Any]]:
    """One row per site the run searched, saying what its search found.

    A value is credited to the search that found it (searched_site_id), wherever it is filed, and
    filed_under_another_site counts how many of them were filed under another of our sites: the
    search looked at another place.
    """
    seed = {_site_id(row): row for row in run.seed}
    accepted: dict[str, list[dict[str, str]]] = defaultdict(list)
    to_review: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in run.resolved:
        if normalize_scalar(row.get("resolution")) in NEW_VALUE_RESOLUTIONS:
            accepted[_searched_id(row)].append(row)
    for row in run.review:
        if normalize_scalar(row.get("resolution")) != CONFIRMED:
            to_review[_searched_id(row)].append(row)
    elsewhere = Counter(
        _searched_id(row) for row in [*run.resolved, *run.review] if _searched_id(row) != _site_id(row)
    )
    decisions: dict[str, Counter] = defaultdict(Counter)
    for row in run.sme or []:
        status = normalize_scalar(row.get("validation_status"))
        if status in ("Validated", "Rejected"):
            decisions[_searched_id(row)][status] += 1
            if status == "Rejected" and normalize_scalar(row.get("rejection_reason")).startswith(WRONG_FACILITY):
                decisions[_searched_id(row)][WRONG_FACILITY] += 1

    rows = []
    for log in run.run_log:
        site_id = _site_id(log)
        site = seed.get(site_id, {})
        name = normalize_scalar(site.get("site_name")) or normalize_scalar(log.get("site_name"))
        status = normalize_scalar(log.get("status"))
        found = [*accepted[site_id], *to_review[site_id]]
        attributes = sorted({normalize_scalar(row.get("attribute_name")) for row in found})
        rows.append({
            "run_id": run.run_id,
            "site_id": site_id,
            "country_iso3": normalize_scalar(site.get("country_iso3")) or normalize_scalar(log.get("country_iso3")),
            "site_name": name,
            "original_site_name": normalize_scalar(site.get("original_site_name")),
            "name_class": name_class(name, site.get("original_site_name")),
            "source_group": source_group(site),
            "status": status,
            "answered": status.startswith(ANSWERED_STATUSES),
            "web_searches": _as_int(log.get("web_searches")),
            "new_values_accepted": len(accepted[site_id]),
            "new_values_to_review": len(to_review[site_id]),
            "found_attributes": "; ".join(attributes),
            "found_name": NAME_ATTRIBUTE in attributes,
            "filed_under_another_site": elsewhere[site_id],
            "not_a_waste_facility": any(
                normalize_scalar(row.get("attribute_name")) == "facility_type"
                and normalize_scalar(row.get("resolved_value")) == NOT_A_WASTE_FACILITY
                for row in found
            ),
            "sme_validated": decisions[site_id]["Validated"],
            "sme_rejected": decisions[site_id]["Rejected"],
            "sme_wrong_facility": decisions[site_id][WRONG_FACILITY],
        })
    return rows


SUMMARY_HEADERS = [
    "source_group", "name_class", "sites_searched", "sites_answered", "bing_queries",
    "queries_per_site", "sites_with_new_value", "new_value_rate", "new_values",
    "queries_per_new_value", "sites_with_found_name", "found_name_rate", "filed_under_another_site",
    "not_a_waste_facility", "sme_validated", "sme_rejected", "sme_wrong_facility", "rate_vs_named",
    "verdict",
]


def _ratio(numerator: float, denominator: float) -> Optional[float]:
    return numerator / denominator if denominator else None


def _stats(rows: list[dict[str, Any]]) -> dict[str, Any]:
    answered = [row for row in rows if row["answered"]]
    queries = sum(row["web_searches"] for row in rows)
    new_values = sum(row["new_values_accepted"] + row["new_values_to_review"] for row in rows)
    with_value = sum(1 for row in answered if row["new_values_accepted"] + row["new_values_to_review"])
    with_name = sum(1 for row in answered if row["found_name"])
    return {
        "sites_searched": len(rows),
        "sites_answered": len(answered),
        "bing_queries": queries,
        "queries_per_site": _ratio(queries, len(rows)),
        "sites_with_new_value": with_value,
        "new_value_rate": _ratio(with_value, len(answered)),
        "new_values": new_values,
        "queries_per_new_value": _ratio(queries, new_values),
        "sites_with_found_name": with_name,
        "found_name_rate": _ratio(with_name, len(answered)),
        "filed_under_another_site": sum(row["filed_under_another_site"] for row in rows),
        "not_a_waste_facility": sum(1 for row in rows if row["not_a_waste_facility"]),
        "sme_validated": sum(row["sme_validated"] for row in rows),
        "sme_rejected": sum(row["sme_rejected"] for row in rows),
        "sme_wrong_facility": sum(row["sme_wrong_facility"] for row in rows),
    }


def verdict(stats: dict[str, Any], named_rate: Optional[float], is_named: bool) -> str:
    """The plan's bar, applied to one class of name within one kind of source."""
    if is_named:
        return "comparison"
    if stats["sites_answered"] < MIN_SITES_TO_JUDGE:
        return "too few to judge"
    decided = stats["sme_validated"] + stats["sme_rejected"]
    if decided and stats["sme_wrong_facility"] * 2 > decided:
        return "below the bar: SMEs rejected most as the wrong facility"
    if not named_rate:
        return "no named sites to compare with"
    if (stats["new_value_rate"] or 0) < BAR * named_rate:
        return "below the bar"
    return "clears the bar"


def summarize(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Per kind of source and class of name, with each class measured against named sites."""
    summary = []
    for group in SOURCE_GROUPS:
        in_group = [row for row in rows if row["source_group"] == group]
        by_class = {cls: _stats([row for row in in_group if row["name_class"] == cls]) for cls in NAME_CLASSES}
        named_rate = by_class[NAMED]["new_value_rate"]
        for cls in NAME_CLASSES:
            stats = by_class[cls]
            if not stats["sites_searched"]:
                continue
            rate = stats["new_value_rate"]
            summary.append({
                "source_group": group,
                "name_class": cls,
                **stats,
                "rate_vs_named": _ratio(rate, named_rate) if rate is not None else None,
                "verdict": verdict(stats, named_rate, cls == NAMED),
            })
    return summary


def queries_per_site(rows: list[dict[str, Any]]) -> tuple[dict[str, float], float]:
    """Measured Bing queries per searched site, by class of name, and over all sites."""
    by_class = {}
    for cls in NAME_CLASSES:
        subset = [row["web_searches"] for row in rows if row["name_class"] == cls]
        if subset:
            by_class[cls] = sum(subset) / len(subset)
    every = [row["web_searches"] for row in rows]
    return by_class, (sum(every) / len(every) if every else 0.0)


FORECAST_HEADERS = ["batch", "name_class", "sites", "est_bing_queries", "est_cost_usd"]


def batch_selections(sites: list[dict[str, Any]]) -> dict[int, list[dict[str, Any]]]:
    """The sites each batch would search without the name rule, selected as a batch run selects
    them (`select_sites`): the default exclusions and sites with nothing left to ask are already
    out, and the classes the rule skips are kept in, since they are what this measures."""
    from .run_context import PipelineConfig
    from .search import select_sites

    selections = {}
    for batch in range(1, LAST_BATCH + 1):
        config = PipelineConfig(
            input_csv=None, run_dir=Path("."), run_id="forecast", batch=batch, include_unsearchable_names=True
        )
        with redirect_stdout(io.StringIO()):  # select_sites narrates; the caller summarizes
            selections[batch] = select_sites(sites, config)
    return selections


SKIP_HEADERS = [
    "batch", "name_class", "reason", "site_id", "site_name", "original_site_name", "source_language",
    "country_iso3", "contributing_data_sources", "latitude", "longitude", "map",
]


def skip_candidates(selections: dict[int, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    """Every site a batch skips by name, with the rule that put it there and a link to look at it."""
    rows = []
    for batch, selected in selections.items():
        for site in selected:
            cls, reason = name_reason(site.get("site_name"), site.get("original_site_name"))
            if cls == NAMED:
                continue
            lat, lon = normalize_scalar(site.get("latitude")), normalize_scalar(site.get("longitude"))
            rows.append({
                "batch": batch,
                "name_class": cls,
                "reason": reason,
                **{key: normalize_scalar(site.get(key)) for key in (
                    "site_id", "site_name", "original_site_name", "source_language", "country_iso3",
                    "contributing_data_sources",
                )},
                "latitude": lat,
                "longitude": lon,
                "map": f"https://www.google.com/maps/search/?api=1&query={lat},{lon}" if lat and lon else "",
            })
    return rows


def format_skip_candidates(selections: dict[int, list[dict[str, Any]]], rows: list[dict[str, Any]], top: int = 15) -> str:
    """Counts by batch and class, then the most common names in each class, to check the rules by eye."""
    lines = ["  batch  searched  " + "  ".join(cls.ljust(13) for cls in NAME_CLASSES[:-1])]
    for batch, selected in selections.items():
        counts = Counter(row["name_class"] for row in rows if row["batch"] == batch)
        lines.append(
            f"  {batch:<5}  {len(selected):>8,}  "
            + "  ".join(f"{counts[cls]:>13,}" for cls in NAME_CLASSES[:-1])
        )
    for cls in NAME_CLASSES[:-1]:
        names = Counter(row["site_name"] for row in rows if row["name_class"] == cls)
        if not names:
            continue
        lines.append(f"\n  Most common {cls} names ({sum(names.values()):,} sites, {len(names):,} different names):")
        lines += [f"    {count:>6,}  {name or '(empty)'}" for name, count in names.most_common(top)]
    return "\n".join(lines)


def forecast(
    sites: list[dict[str, Any]],
    per_site: dict[str, float],
    fallback: float,
    cost_per_1000_queries: float = DEFAULT_COST_PER_1000_QUERIES,
) -> list[dict[str, Any]]:
    """What each batch would search without the name rule, by class of name, and what that would
    cost: the classes the rule skips show what it saves.

    Sites are selected as a batch run selects them (`select_sites`), so the default exclusions and
    sites with nothing left to ask are already out.
    """
    rows = []
    for batch, selected in batch_selections(sites).items():
        counts = Counter(name_class(site.get("site_name"), site.get("original_site_name")) for site in selected)
        for cls in NAME_CLASSES:
            queries = counts[cls] * per_site.get(cls, fallback)
            rows.append({
                "batch": batch,
                "name_class": cls,
                "sites": counts[cls],
                "est_bing_queries": round(queries),
                "est_cost_usd": round(queries * cost_per_1000_queries / 1000, 2),
            })
    return rows


def describe_run(run: RunFiles) -> str:
    sme = "no reviewed workbook found" if run.sme is None else f"{len(run.sme):,} reviewed-workbook rows"
    return (
        f"Run {run.run_id}: {len(run.run_log):,} sites searched; {sme}.\n"
        f"  Consolidation run its seed read: {describe_manifest_run(run.manifest)}."
    )


def _show(value: Any, percent: bool = False) -> str:
    if value is None:
        return "-"
    if percent:
        return f"{value:.0%}"
    if isinstance(value, float):
        return f"{value:.1f}"
    return f"{value:,}"


def format_summary(summary: list[dict[str, Any]]) -> str:
    """The summary as a plain-text table, one block per kind of source."""
    columns = [
        ("class", lambda r: r["name_class"], 14),
        ("searched", lambda r: _show(r["sites_searched"]), 9),
        ("queries/site", lambda r: _show(r["queries_per_site"]), 13),
        ("new value", lambda r: _show(r["new_value_rate"], percent=True), 10),
        ("found name", lambda r: _show(r["found_name_rate"], percent=True), 11),
        ("filed elsewhere", lambda r: _show(r["filed_under_another_site"]), 16),
        ("SME ok/rej/wrong", lambda r: f"{r['sme_validated']}/{r['sme_rejected']}/{r['sme_wrong_facility']}", 17),
        ("vs named", lambda r: _show(r["rate_vs_named"]), 9),
        ("verdict", lambda r: r["verdict"], 0),
    ]
    lines = []
    for group in SOURCE_GROUPS:
        block = [row for row in summary if row["source_group"] == group]
        if not block:
            continue
        lines.append(f"\nSites from {group}" if group == MAP_ONLY else "\nSites with another source")
        lines.append("  " + "".join(title.ljust(width) for title, _, width in columns))
        for row in block:
            lines.append("  " + "".join(str(cell(row)).ljust(width) for _, cell, width in columns))
    return "\n".join(lines)


def write_csv(path: Path, rows: list[dict[str, Any]], headers: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
